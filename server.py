from __future__ import annotations

import json
import hashlib
import fcntl
import re
import socket
import tempfile
import urllib.error
import urllib.request
try:
    import cgi
except ModuleNotFoundError:
    cgi = None
from email.parser import BytesParser
from email.policy import default as email_default_policy
from html import unescape
import io
import os
import queue
import shutil
import subprocess
import threading
import uuid
from contextlib import contextmanager
from copy import deepcopy
from datetime import date, datetime, timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from importlib.util import module_from_spec, spec_from_file_location
from pathlib import Path
from urllib.parse import parse_qs, quote, unquote, urlparse
from http.cookies import CookieError, SimpleCookie

from docx import Document
from dotenv import load_dotenv

from app_storage import AppStore, IdempotencyConflict, ProjectVersionConflict
from mail_center import (
    normalize_smtp_settings,
    parse_recipient_file,
    render_mail_text,
    send_smtp_message,
    test_smtp_connection,
    validate_mail_compose,
)
from project_domain import (
    build_bpm_preflight,
    normalize_project_state,
    project_state_from_report,
    summarize_project_state,
)
from project_files import InvalidProjectFile, ProjectFileStore


ROOT = Path(__file__).resolve().parent
load_dotenv(ROOT / ".env")
TEMPLATE = ROOT / "templates" / "planning-report-template.docx"
EXTRACTOR = ROOT / "skills" / "phei-topic-planning-report" / "scripts" / "extract_application_docx.py"
REPORT_RULES = ROOT / "skills" / "phei-topic-planning-report" / "references" / "report_sections.md"
SCORING_RULES = ROOT / "skills" / "phei-topic-planning-report" / "references" / "scoring_rules.md"
BPM_SCRIPT = ROOT / "skills" / "phei-bpm-topic-declaration" / "scripts" / "fill_topic.js"
NODE_MODULES = ROOT / "node_modules"
RUNTIME_LOG = ROOT / "server-runtime.log"

APP_STORE = AppStore(
    Path(os.environ.get("PHEI_DB_PATH", ROOT / "data" / "app.db")),
    os.environ.get("APP_CREDENTIAL_KEY", ""),
)
PROJECT_FILE_STORE = ProjectFileStore(
    Path(os.environ.get("PHEI_PROJECTS_PATH", ROOT / "data" / "projects"))
)
SESSION_COOKIE_NAME = "phei_session"
STATIC_ASSETS = {
    "/": ("index.html", "text/html; charset=utf-8"),
    "/index.html": ("index.html", "text/html; charset=utf-8"),
    "/auth_session_generation.js": ("auth_session_generation.js", "application/javascript; charset=utf-8"),
    "/mail_ui_helpers.js": ("mail_ui_helpers.js", "application/javascript; charset=utf-8"),
    "/vendor/lucide-0.468.0.min.js": (
        "vendor/lucide-0.468.0.min.js",
        "application/javascript; charset=utf-8",
    ),
    "/assets/report-preview.png": ("assets/report-preview.png", "image/png"),
}

PROJECT_API_RE = re.compile(
    r"^/api/projects/([0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12})"
    r"(?:/(preflight|archive|files)(?:/([A-Za-z0-9_-]{1,128}))?)?$"
)
MAIL_TEMPLATE_API_RE = re.compile(r"^/api/mail/templates/([0-9a-f]{12})$")
MAIL_BATCH_API_RE = re.compile(r"^/api/mail/batches/([0-9a-f]{12})(?:/(retry))?$")
DEFAULT_MAIL_TEMPLATES = [
    {
        "name": "教材合作邀请",
        "subject": "诚邀您参与《教材名称》教材建设",
        "body": "{{姓名}}老师，您好！\n\n我们正在开展教材建设工作，诚邀您参与。\n\n此致\n敬礼！",
    },
    {
        "name": "材料提醒",
        "subject": "《教材名称》材料提交提醒",
        "body": "{{姓名}}老师，您好！\n\n烦请您在方便时查看并提交相关材料。\n\n谢谢！",
    },
]
MAX_MAIL_RECIPIENTS = 1000
MAX_MAIL_RECIPIENT_FILE_BYTES = 10 * 1024 * 1024
MAX_JSON_BODY_BYTES = 2 * 1024 * 1024
JSON_READ_TIMEOUT_SECONDS = 15
IDEMPOTENCY_KEY_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9._:-]{7,127}\Z")
SERVER_LOCK_PATH = Path(
    os.environ.get("PHEI_LOCK_PATH", str(APP_STORE.db_path.with_suffix(".server.lock")))
)


class RequestHandled(Exception):
    pass


class MultipartForm:
    def __init__(self, fp, headers, environ):
        content_type = headers.get("Content-Type", "")
        content_length = int(environ.get("CONTENT_LENGTH", "0"))
        raw_message = (
            f"Content-Type: {content_type}\r\nMIME-Version: 1.0\r\n\r\n".encode("utf-8")
            + fp.read(content_length)
        )
        message = BytesParser(policy=email_default_policy).parsebytes(raw_message)
        self.fields = {}
        for part in message.iter_parts():
            name = part.get_param("name", header="content-disposition")
            if not name or name in self.fields:
                continue
            content = part.get_payload(decode=True) or b""
            filename = part.get_filename()
            if filename is None:
                self.fields[name] = content.decode(part.get_content_charset() or "utf-8")
            else:
                self.fields[name] = type(
                    "UploadedFile", (), {"file": io.BytesIO(content), "filename": filename}
                )()

    def __contains__(self, name):
        return name in self.fields

    def __getitem__(self, name):
        return self.fields[name]

    def getfirst(self, name, default=None):
        return self.fields.get(name, default)


def read_multipart_form(fp, headers, environ):
    if cgi is not None:
        return cgi.FieldStorage(fp=fp, headers=headers, environ=environ)
    return MultipartForm(fp=fp, headers=headers, environ=environ)

DEFAULT_MODEL_URL = os.environ.get("DEEPSEEK_BASE_URL", "https://api.deepseek.com").strip()
DEFAULT_API_KEY = os.environ.get("DEEPSEEK_API_KEY", "").strip()
DEFAULT_MODEL = os.environ.get("DEEPSEEK_MODEL", "deepseek-v4-pro").strip()
FAST_MODEL = os.environ.get("DEEPSEEK_FAST_MODEL", "deepseek-v4-flash").strip()

MODEL_CONFIGURATION_ERROR = "模型服务暂不可用，请联系管理员检查服务端配置。"
MODEL_REQUEST_ERROR = "模型服务暂不可用，请稍后重试。"


class ModelServiceError(RuntimeError):
    pass


class ModelConfigurationError(ModelServiceError, ValueError):
    pass

JOB_QUEUE: queue.Queue = queue.Queue()


def runtime_log(message: str):
    try:
        stamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with RUNTIME_LOG.open("a", encoding="utf-8") as log:
            log.write(f"{stamp} {message}\n")
    except Exception:
        pass


ROW_BY_KEY = {
    "content": 7,
    "author": 8,
    "feasibility": 9,
    "award": 10,
    "profit": 11,
    "marketing": 12,
}

SECTION_TITLES = {
    "content": "一、选题内容",
    "author": "二、作者情况",
    "feasibility": "三、策划过程与可行性",
    "award": "四、获奖潜质",
    "profit": "五、成本与盈利估算",
    "marketing": "六、市场定位与营销",
}

SCORE_NAMES = {
    "content": "选题内容",
    "author": "作者情况",
    "feasibility": "策划过程与可行性",
    "award": "获奖潜质",
    "profit": "成本与盈利估算",
    "marketing": "市场定位与营销",
}

BPM_FIXED_CLASSIFICATION = {
    "class1": "02",
    "class2": "0201",
    "gbClass": "G",
    "readLevel": "高等理工",
}

FIXED_PROFIT_SECTION_TEXT = """　　（一）纸质教材
　　1.出版规格
　　（1）版权字数： 350 千字。
　　（2）印张： 14。
　　（3）正文印刷色数：  单色。

　　2.盈利估算
　　（1）定价：  59.8  元。
　　（2）版税（率）：  8%。
　　（3）发货折扣：    65%。
　　（4）预估首印数：   1200册。
　　（5）预估总印数：   3000册。
　　（6）包销册数：     0册；包销折扣： 0 %。
　　（7）出版资助：   
　　（8）预计总毛利润：   """

AI_REPORT_STYLE_GUIDE = """参照《AI驱动软件开发实战》选题策划报告的一到六部分写法，对除第五部分外的内容进行二次润色。报告供出版社领导审议，目标是以真实事实论证选题的出版价值和实施可行性，帮助领导判断是否批准立项。风格要求如下：
1. 一、选题内容：写成3段左右。第1段判断出版导向、现实需求和出版定位；第2段概括全书规模、篇章结构和主要内容链条；第3段集中写创新点和差异化价值。版权情况仅在来源有依据时表述，没有依据则直接略去。
2. 二、作者情况：写成2段左右。第1段写第一作者身份、职业/教学/科研经历和与选题的相关性；第2段写来源明确的代表性成果、项目、著作、专利、课程或团队基础。只论证作者完成本书的能力，不写给责任编辑或作者的工作建议。
3. 三、策划过程与可行性：写成3段左右。第1段从真实教学、行业或课程痛点切入；第2段写策划思路，强调“概念-方法-案例-工具包”或类似闭环；第3段依据已有稿件、目录、进度和资源素材，直接论证按期出版可行性，不布置组稿、跟进或核实任务。
4. 四、获奖潜质：写成2段左右。依据已有教材定位、内容创新、课程应用和资源建设基础，客观说明成果培育、资源传播或版权输出潜力；语气谨慎，不承诺获奖。没有具体获奖材料时，不写“缺少获奖依据”等缺项说明。
5. 六、市场定位与营销：写成4段左右。第1段列出2-4类目标读者；第2段写竞品/同类选题格局和本书差异化；第3段以“本书可通过”“推广将围绕”等策划表述写课程推广、作者渠道、技术/教学社区、样章试读、讲座培训等；第4段写来源支持的配套资源和“图书+资源+服务”价值，不向责任编辑布置操作任务。
6. 整体语言应具体、稳健、积极、有判断。每段都要回答“为什么值得出版、内容有什么价值、现有基础为什么能够支撑落地”，不得写成风险提示、内部工作安排或审批后的操作建议。
7. 只写申报表能够支持的事实和合理判断。某项信息缺失时直接略去，不得在一到六正文中写“申报表未提供”“申报表中没有体现”“目前尚不明确”“缺少相关依据”“有待进一步确认”等缺项说明。
8. 不得出现“建议责任编辑”“建议在组稿阶段”“建议后续跟进”“建议进一步确认”“不建议在论证中强调”“待作者补充后再”等内部工作指令。
9. 不要照抄范文中的AI软件开发事实，必须替换为当前申报表对应事实。
"""

REPORT_TONE_BANNED_PATTERNS = (
    r"建议(?:责任编辑|策划编辑|作者)",
    r"建议(?:在)?组稿(?:阶段)?",
    r"建议(?:后续|进一步)(?:跟进|确认|核实|补充|完善|商议)",
    r"不建议.{0,12}(?:强调|写入|表述|宣传)",
    r"待作者.{0,12}(?:补充|确认|提供|完善)",
    r"(?:申报表|申报材料|材料)(?:中)?(?:未|没有)(?:提供|体现|说明|列明|明确|填写)?",
    r"(?:申报表|申报材料|材料)(?:中)?(?:缺少|缺乏)",
    r"(?:作者|团队|本书|本选题).{0,10}(?:缺少|缺乏)(?:国家级|省部级|明确|具体|相关)?",
    r"缺少(?:明确|具体|相关)?(?:奖项|依据|材料|信息|记录|支撑|成果|数据|说明)",
    r"(?:目前|当前)?尚未(?:明确|体现|提供|确认)",
    r"有待进一步确认",
)


def safe_name(name: str) -> str:
    name = re.sub(r"[\\/:*?\"<>|]+", "_", name or "选题策划报告")
    return name.strip() or "选题策划报告"


def format_chinese_date(value: date | None = None) -> str:
    value = value or date.today()
    return f"{value.year}年{value.month}月{value.day}日"


def editor_name_from_payload(payload: dict) -> str:
    bpm = payload.get("bpm") or {}
    return first_non_empty(payload.get("editorName", ""), bpm.get("name", ""), "叶文涛")


def title_for_report_filename(title: str) -> str:
    clean = safe_name(title or "选题策划报告")
    if clean.startswith("《") and clean.endswith("》"):
        return clean
    return f"《{clean}》"


def report_file_stem(title: str, editor_name: str, value: date | None = None) -> str:
    return f"{title_for_report_filename(title)}选题策划报告（高等教育出版分社+{safe_name(editor_name)}+{format_chinese_date(value)}）"


def first_unique_cell(row, index):
    seen = []
    unique = []
    for cell in row.cells:
        if cell._tc not in seen:
            seen.append(cell._tc)
            unique.append(cell)
    return unique[index]


def capture_format(cell):
    para = cell.paragraphs[0] if cell.paragraphs else None
    ppr = deepcopy(para._p.pPr) if para is not None and para._p.pPr is not None else None
    rpr = None
    if para is not None:
        for run in para.runs:
            if run._r.rPr is not None:
                rpr = deepcopy(run._r.rPr)
                break
    return ppr, rpr


def ensure_indent(text: str) -> str:
    parts = []
    for part in re.split(r"\n\s*\n", text or ""):
        lines = [line.rstrip() for line in part.splitlines()]
        cleaned = "\n".join(lines).strip()
        if not cleaned:
            continue
        cleaned = re.sub(r"^[　\s]+", "", cleaned)
        parts.append("　　" + cleaned)
    return "\n\n".join(parts)


def set_cell_text_preserve_format(cell, text: str):
    ppr, rpr = capture_format(cell)
    cell._tc.clear_content()
    parts = [part for part in re.split(r"\n\s*\n", text or "") if part.strip()] or [""]
    for part in parts:
        para = cell.add_paragraph()
        if ppr is not None:
            para._p.insert(0, deepcopy(ppr))
        for i, line in enumerate(part.splitlines() or [""]):
            if i:
                para.add_run().add_break()
            run = para.add_run(line)
            if rpr is not None:
                run._r.insert(0, deepcopy(rpr))


def clear_fixed_row_height(row):
    tr_pr = row._tr.trPr
    if tr_pr is None:
        return
    for child in list(tr_pr):
        if child.tag.endswith("}trHeight"):
            tr_pr.remove(child)


def build_docx(payload: dict) -> Path:
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"模板不存在: {TEMPLATE}")

    raw_title = payload.get("title") or "选题策划报告"
    title = safe_name(raw_title)
    editor_name = editor_name_from_payload(payload)
    today = date.today()
    out = Path(tempfile.gettempdir()) / f"{report_file_stem(raw_title, editor_name, today)}.docx"
    doc = Document(str(TEMPLATE))
    table = doc.tables[0]

    set_cell_text_preserve_format(first_unique_cell(table.rows[0], 1), format_chinese_date(today))
    set_cell_text_preserve_format(first_unique_cell(table.rows[2], 1), title)
    set_cell_text_preserve_format(first_unique_cell(table.rows[3], 1), editor_name)
    set_cell_text_preserve_format(first_unique_cell(table.rows[4], 5), str(payload.get("total", "")))
    set_cell_text_preserve_format(first_unique_cell(table.rows[5], 5), str(payload.get("total", "")))

    scores = {item[0]: item[2] for item in payload.get("scores", []) if len(item) >= 3}
    score_by_key = {
        "content": scores.get("选题内容", ""),
        "author": scores.get("作者情况", ""),
        "feasibility": scores.get("策划过程与可行性", ""),
        "award": scores.get("获奖潜质", ""),
        "profit": scores.get("成本与盈利估算", ""),
        "marketing": scores.get("市场定位与营销", ""),
    }

    for section in payload.get("sections", []):
        key = section.get("key")
        if key not in ROW_BY_KEY:
            continue
        row_idx = ROW_BY_KEY[key]
        set_cell_text_preserve_format(first_unique_cell(table.rows[row_idx], 2), ensure_indent(section.get("text", "")))
        set_cell_text_preserve_format(first_unique_cell(table.rows[row_idx], 3), str(score_by_key.get(key, "")))

    clear_fixed_row_height(table.rows[13])
    set_cell_text_preserve_format(first_unique_cell(table.rows[13], 2), "本选题如为新编教材，则此项不适用；如为再版修订，请补充上一版销量、院校使用情况和修订计划。")
    set_cell_text_preserve_format(first_unique_cell(table.rows[14], 2), "目录待由申报表解析后写入。")
    set_cell_text_preserve_format(first_unique_cell(table.rows[15], 2), "样章内容待作者补充。")

    doc.save(str(out))
    return out


def load_extractor():
    spec = spec_from_file_location("phei_extract_application_docx", EXTRACTOR)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法加载申报表抽取脚本: {EXTRACTOR}")
    module = module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8")


AUTHOR_BIO_FIELDS = ("作者简介", "作译者简介", "主编简介", "工作简历或单位简介", "工作简历", "个人简介")


def html_to_text(html: str) -> str:
    html = re.sub(r"(?is)<(script|style|noscript).*?>.*?</\1>", " ", html or "")
    html = re.sub(r"(?is)<br\s*/?>", "\n", html)
    html = re.sub(r"(?is)</p\s*>", "\n", html)
    html = re.sub(r"(?is)<[^>]+>", " ", html)
    text = unescape(html)
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    return re.sub(r"\n\s*\n+", "\n", text).strip()


def fetch_url_text(url: str, timeout: int = 6, limit: int = 160000) -> str:
    try:
        request = urllib.request.Request(
            url,
            headers={
                "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 Chrome/125 Safari/537.36",
                "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.6",
            },
        )
        with urllib.request.urlopen(request, timeout=timeout) as response:
            raw = response.read(limit)
            charset = response.headers.get_content_charset() or "utf-8"
        return raw.decode(charset, errors="ignore")
    except Exception as exc:
        runtime_log(f"author-search fetch failed url={url} error={type(exc).__name__}: {exc}")
        return ""


def parse_search_results(html: str) -> list[dict]:
    results = []
    for match in re.finditer(r'(?is)<li[^>]+class="[^"]*\bb_algo\b[^"]*"[^>]*>.*?<h2[^>]*>\s*<a[^>]+href="([^"]+)"[^>]*>(.*?)</a>.*?</li>', html or ""):
        url = unescape(match.group(1))
        title = html_to_text(match.group(2))
        if url.startswith("http"):
            results.append({"title": title, "url": url})
    if results:
        return results[:8]
    for match in re.finditer(r'(?is)<a[^>]+href="(https?://[^"]+)"[^>]*>(.*?)</a>', html or ""):
        url = unescape(match.group(1))
        title = html_to_text(match.group(2))
        if title and "bing" not in urlparse(url).netloc:
            results.append({"title": title, "url": url})
        if len(results) >= 8:
            break
    return results


def parse_exa_search_output(output: str) -> list[dict]:
    results = []
    for block in re.split(r"\n\s*---\s*\n", output or ""):
        title_match = re.search(r"(?m)^Title:\s*(.+)$", block)
        url_match = re.search(r"(?m)^URL:\s*(https?://\S+)$", block)
        if not title_match or not url_match:
            continue
        highlight_match = re.search(r"(?s)Highlights:\s*(.+)$", block)
        results.append(
            {
                "title": strip_indent(title_match.group(1)),
                "url": strip_indent(url_match.group(1)),
                "excerpt": short_text(
                    strip_indent(highlight_match.group(1)) if highlight_match else "",
                    2400,
                ),
            }
        )
    return results[:8]


def search_web_results_exa(query: str, limit: int = 8) -> list[dict]:
    executable = shutil.which("mcporter")
    if not executable:
        return []
    expression = (
        "exa.web_search_exa(query: "
        f"{json.dumps(query, ensure_ascii=False)}, numResults: {max(1, min(8, limit))})"
    )
    try:
        completed = subprocess.run(
            [executable, "call", expression],
            capture_output=True,
            text=True,
            timeout=25,
            check=False,
        )
    except (OSError, subprocess.SubprocessError) as exc:
        runtime_log(f"exa search failed error={type(exc).__name__}: {exc}")
        return []
    if completed.returncode != 0:
        runtime_log(f"exa search failed status={completed.returncode}")
        return []
    return parse_exa_search_output(completed.stdout)


def search_web_results(query: str) -> list[dict]:
    exa_results = search_web_results_exa(query)
    if exa_results:
        return exa_results
    url = f"https://www.bing.com/search?q={quote(query)}&setlang=zh-CN"
    html = fetch_url_text(url, timeout=7, limit=220000)
    return parse_search_results(html)


def official_score(result: dict, unit: str) -> int:
    url = result.get("url", "")
    host = urlparse(url).netloc.lower()
    text = f"{result.get('title', '')} {url}".lower()
    score = 0
    if ".edu.cn" in host or ".edu." in host or host.endswith(".edu"):
        score += 5
    if "faculty" in text or "teacher" in text or "szdw" in text or "师资" in text or "教师" in text:
        score += 2
    unit_tokens = [token for token in re.split(r"[\s,，/／（）()·]+", unit or "") if len(token) >= 2]
    if any(token in result.get("title", "") or token in url for token in unit_tokens[:3]):
        score += 2
    return score


def author_official_search_context(facts: dict) -> dict:
    if first_non_empty(*(field_value(facts, name) for name in AUTHOR_BIO_FIELDS)):
        return {}
    name = first_non_empty(
        field_value(facts, "作者姓名", "姓名", "主要作（译）者姓名", "第一作者姓名"),
        field_value(facts, "主编", "作者"),
    )
    unit = field_value(facts, "工作单位", "单位名称", "作者单位", "所在单位")
    if not name:
        return {}
    query = " ".join(part for part in [name, unit, "教师 简介 个人主页"] if part)
    try:
        results = search_web_results(query)
    except Exception as exc:
        runtime_log(f"author-search query failed name={name} error={type(exc).__name__}: {exc}")
        return {}
    ranked = sorted(results, key=lambda item: official_score(item, unit), reverse=True)
    selected = []
    for result in ranked[:4]:
        page_text = strip_indent(str(result.get("excerpt", "") or ""))
        if not page_text:
            page_html = fetch_url_text(result["url"], timeout=5, limit=120000)
            page_text = html_to_text(page_html)
        if name not in page_text and official_score(result, unit) < 5:
            excerpt = ""
        else:
            index = page_text.find(name)
            if index >= 0:
                start = max(0, index - 240)
                excerpt = page_text[start : start + 1200]
            else:
                excerpt = page_text[:900]
        selected.append(
            {
                "title": result.get("title", ""),
                "url": result.get("url", ""),
                "excerpt": short_text(excerpt, 900),
            }
        )
        if len(selected) >= 3:
            break
    context = {
        "query": query,
        "results": [item for item in selected if item.get("excerpt")],
    }
    if context["results"]:
        runtime_log(f"author-search ok name={name} results={len(context['results'])}")
    else:
        runtime_log(f"author-search no usable official text name={name}")
    return context if context["results"] else {}


BOOK_SOURCE_HOST_SCORES = {
    "ecsponline.com": 6,
    "cepp.sgcc.com.cn": 6,
    "pup.cn": 6,
    "tup.tsinghua.edu.cn": 6,
    "epubit.com": 6,
    "cmpbook.com": 6,
    "dushu.com": 3,
    "jd.com": 2,
    "dangdang.com": 2,
}


def comparison_topic(title: str) -> str:
    topic = re.sub(
        r"(?:项目式|案例式|新形态|数字|一体化)?(?:教程|教材|实战|基础与应用|技术及应用|应用技术)$",
        "",
        strip_indent(title),
    )
    return topic.strip(" ：:—-《》") or strip_indent(title)


def book_title_identity(title: str) -> str:
    identity = re.split(r"_|\s+-\s+|[|｜]", strip_indent(title), maxsplit=1)[0]
    return re.sub(r"[《》\s]", "", identity).lower()


def book_search_score(result: dict, topic: str) -> int:
    url = str(result.get("url", "") or "")
    title = str(result.get("title", "") or "")
    host = urlparse(url).netloc.lower()
    score = max(
        (value for domain, value in BOOK_SOURCE_HOST_SCORES.items() if domain in host),
        default=0,
    )
    if any(token in title for token in ("出版社", "图书", "教材", "书店", "商城")):
        score += 2
    topic_tokens = [token for token in re.split(r"[与及和、：:\s]+", topic) if len(token) >= 2]
    score += sum(1 for token in topic_tokens[:4] if token in title)
    if any(token in title for token in ("论文", "会议", "招聘", "新闻", "政策")):
        score -= 4
    return score


def book_comparison_search_context(facts: dict) -> list[dict]:
    title = field_value(facts, "教材名称", "选题名称", "选题名", "书名")
    if not title:
        return []
    topic = comparison_topic(title)
    queries = [
        f"{topic} 同类图书 教材 出版社",
        f"{topic} 专著 图书 ISBN 目录",
    ]
    candidates = []
    seen_urls = set()
    for query in queries:
        try:
            results = search_web_results(query)
        except Exception as exc:
            runtime_log(f"book-search query failed title={title} error={type(exc).__name__}: {exc}")
            continue
        for result in results:
            url = str(result.get("url", "") or "")
            if not url.startswith("http") or url in seen_urls:
                continue
            seen_urls.add(url)
            score = book_search_score(result, topic)
            if score < 2:
                continue
            candidates.append((score, result))

    selected = []
    seen_books = set()
    for score, result in sorted(candidates, key=lambda item: item[0], reverse=True)[:8]:
        identity = book_title_identity(str(result.get("title", "") or ""))
        if not identity or identity in seen_books:
            continue
        page_text = strip_indent(str(result.get("excerpt", "") or ""))
        if not page_text:
            page_html = fetch_url_text(result["url"], timeout=6, limit=180000)
            page_text = html_to_text(page_html)
        if not page_text:
            continue
        if not any(marker in page_text for marker in ("出版社", "ISBN", "书号", "目录", "内容简介")):
            continue
        selected.append(
            {
                "title": short_text(result.get("title", ""), 160),
                "url": result.get("url", ""),
                "excerpt": short_text(page_text, 1800),
                "sourceScore": score,
            }
        )
        seen_books.add(identity)
        if len(selected) >= 6:
            break
    runtime_log(f"book-search complete title={title} candidates={len(selected)}")
    return selected


def compact_facts(facts: dict) -> dict:
    fields = facts.get("canonical_fields") or facts.get("application_fields") or {}
    keep_keys = [
        "教材名称",
        "选题名",
        "选题名称",
        "作者姓名",
        "姓名",
        "单位名称",
        "职务/职称",
        "职称",
        "学术或教育组织任职",
        "已出版教材或著作",
        "作者简介",
        "作译者简介",
        "工作单位",
        "作者单位",
        "性别",
        "证件类型",
        "证件号",
        "身份证号",
        "学历",
        "学位",
        "专业",
        "作者专业",
        "毕业院校",
        "联系电话",
        "联系电话1",
        "联系电话2",
        "手机",
        "电子邮箱",
        "邮箱",
        "邮编",
        "通信地址",
        "通讯地址",
        "单位地址",
        "单位邮编",
        "单位联系人",
        "工作简历或单位简介",
        "合作者 情况简介",
        "合作者情况简介",
        "参加的学术组织及任职情况",
        "科研或教研项目经历",
        "科研或教学工作及获奖情况",
        "著作方向",
        "主要著作出版情况",
        "近几年 教学经历",
        "本课程教改、效果和获奖情况",
        "内容简介",
        "读者 对象",
        "读者定位",
        "本教材特色和优势",
        "国内外同类教材比较",
        "课程基本情况",
        "估计字数",
        "估计字数（千字）",
        "适用层次",
        "适用专业",
        "类别",
        "总学时数",
        "本校教材年用量估计",
        "是否愿意资助出版，资助费用",
        "预定交稿时间",
        "著译写作 时间表",
        "参编人员",
        "编 写 大 纲 或 主 要 目 录",
        "写作资料 主要来源",
    ]
    selected = {key: fields.get(key) for key in keep_keys if fields.get(key)}
    compact = {
        "document_title": facts.get("document_title"),
        "form_date": facts.get("form_date"),
        "fields": selected,
        "appendix_sections": facts.get("appendix_sections", {}),
    }
    official_context = author_official_search_context(facts)
    if official_context:
        compact["author_official_search_context"] = official_context
    return compact


def strip_indent(text: str) -> str:
    text = re.sub(r"^[　\s]+", "", text or "", flags=re.M)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def field_value(facts: dict, *names: str) -> str:
    fields = facts.get("canonical_fields") or facts.get("application_fields") or {}
    for name in names:
        value = fields.get(name)
        if value:
            return str(value).strip()
    return ""


def section_value(result: dict, key: str) -> str:
    for section in result.get("sections", []):
        if section.get("key") == key:
            return strip_indent(section.get("text", ""))
    return ""


def first_non_empty(*values: str) -> str:
    for value in values:
        clean = strip_indent(value)
        if clean:
            return clean
    return ""


def clean_person_name(value: str) -> str:
    text = strip_indent(value)
    if not text:
        return ""
    text = re.sub(r"^(作者姓名|姓名|第一作者姓名|主要作（译）者姓名|主编|作者)[:：\s]*", "", text)
    text = text.strip(" ：:，,；;。")
    if is_likely_person_name(text):
        return text
    patterns = [
        r"第一作者\s*([\u4e00-\u9fff·]{2,8})",
        r"本书第一作者\s*([\u4e00-\u9fff·]{2,8})",
        r"作者[:：\s]+([\u4e00-\u9fff·]{2,8})",
        r"主编[:：\s]+([\u4e00-\u9fff·]{2,8})",
        r"姓名[:：\s]*([\u4e00-\u9fff·]{2,8})",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match and is_likely_person_name(match.group(1).strip()):
            return match.group(1).strip()
    first = re.split(r"[，,、；;\s\n]", text, maxsplit=1)[0].strip()
    if is_likely_person_name(first):
        return first
    return ""


def is_likely_person_name(value: str) -> bool:
    text = strip_indent(value).strip(" ：:，,；;。")
    if not re.fullmatch(r"[\u4e00-\u9fff·]{2,8}", text or ""):
        return False
    if text.startswith(("为", "由", "以", "与", "和", "及", "等")):
        return False
    non_name_words = (
        "自由职业", "职业者", "第一作者", "合作者", "作者", "主编", "教材",
        "本书", "该书", "出版", "情况", "简介", "工程师", "教授", "研究员",
    )
    return not any(word in text for word in non_name_words)


def all_fact_fields(facts: dict) -> dict:
    fields = {}
    for source_name in ("application_fields", "canonical_fields"):
        source = facts.get(source_name) or {}
        if isinstance(source, dict):
            fields.update(source)
    return fields


def extract_first_author_from_text(value: str) -> str:
    text = strip_indent(value)
    if not text:
        return ""
    first = re.split(r"[，,、；;\s\n]", text, maxsplit=1)[0].strip()
    if is_likely_person_name(first):
        return first
    patterns = [
        r"([\u4e00-\u9fff·]{2,8})[，,、]\s*(?:第一作者|主编|作者|主笔)",
        r"(?:第一作者|主编|作者|姓名)[:：\s]+([\u4e00-\u9fff·]{2,8})",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match and is_likely_person_name(match.group(1).strip()):
            return match.group(1).strip()
    return ""


def extract_author_name_from_facts(facts: dict) -> str:
    for key in ("作者姓名", "姓名", "主要作（译）者姓名", "第一作者姓名", "主编"):
        name = clean_person_name(field_value(facts, key))
        if name:
            return name

    for key, value in all_fact_fields(facts).items():
        key_text = str(key)
        if any(token in key_text for token in ("姓名", "作者", "合作者", "主编")):
            name = extract_first_author_from_text(str(value))
            if name:
                return name
    return ""


def blank_if_missing(value: str) -> str:
    return strip_indent(value) or "　"


def short_text(text: str, limit: int = 1000) -> str:
    clean = strip_indent(text)
    return clean if len(clean) <= limit else clean[: limit - 1] + "…"


def number_text(value: str, default: str) -> str:
    match = re.search(r"\d+(?:\.\d+)?", str(value or ""))
    return match.group(0) if match else default


def normalize_bpm_date(value: str, fallback: date) -> str:
    text = str(value or "")
    match = re.search(r"(20\d{2})[年./-]\s*(\d{1,2})(?:[月./-]\s*(\d{1,2}))?", text)
    if not match:
        return fallback.isoformat()
    year = int(match.group(1))
    month = int(match.group(2))
    day = int(match.group(3) or 28)
    try:
        return date(year, month, day).isoformat()
    except ValueError:
        return fallback.isoformat()


def infer_bpm_detailed_classification(title: str, facts: dict, content: str, marketing: str) -> dict:
    text = " ".join(
        [
            title or "",
            field_value(facts, "适用专业", "类别", "课程基本情况"),
            content or "",
            marketing or "",
        ]
    )
    if re.search(r"人工智能|AI|软件|程序|编程|Python|Java|Web|数据|数据库|算法|计算机", text, re.I):
        return {
            "class3": "020101",
            "class4": "02010103",
            "bwClass": "计算机",
            "bwCipClass": "TP",
        }
    if re.search(r"经济|管理|营销|财务|会计|金融|工商", text):
        return {
            "class3": "",
            "class4": "",
            "bwClass": "经管",
            "bwCipClass": "F",
        }
    if re.search(r"物理|医学|医药|机械|电子|通信|自动化|工程|材料", text):
        return {
            "class3": "020101",
            "class4": "02010103",
            "bwClass": "大科技专业与科普",
            "bwCipClass": "N",
        }
    return {
        "class3": "020101",
        "class4": "02010103",
        "bwClass": "其他",
        "bwCipClass": "G",
    }


def fixed_bpm_defaults() -> dict:
    return {
        "readerNum": "100",
        "language": "中文",
        "scriptSource": "作者独立投稿",
        "scriptStyle": "电子文件",
        "awards": "无",
        "colorPrint": "单色",
        "haveCd": "",
        "words": "350.00",
        "price": "59.8",
        "remCharNum": "0",
        "remPayMode": "销数版税",
        "remStandard": "8",
        "remUnit": "2",
        "publishMode": "常规出版",
        "haveZz": "否",
        "imburseFee": "0",
        "imburseNum": "0",
        "haveBx": "否",
        "bsaleNum": "0",
        "bsaleDiscount": "0",
        "coMode": "",
        "coBuyDiscount": "0",
        "bsaleMode": "",
        "digital": "授权",
        "eRemPayMode": "是",
        "eRemStandard": "8",
        "totalNum": "3000",
        "firstNum": "1200",
        "project": "无",
        "scriptClassify": "普通选题",
        "isCost": "1",
        "impScript": "否",
        "isUrgent": "普通",
        "isMeeting": "否",
    }


def author_field(facts: dict, *names: str) -> str:
    return field_value(facts, *names)


def build_author_bio(facts: dict, result: dict) -> tuple[str, str]:
    source_bio = first_non_empty(
        author_field(facts, "作者简介", "作译者简介", "主编简介"),
        author_field(facts, "工作简历或单位简介", "工作简历", "个人简介"),
    )
    if source_bio:
        return short_text(source_bio, 1000), "申报表已有"

    model_author = ""
    author_maintenance = result.get("authorMaintenance") or result.get("author_maintenance") or {}
    if isinstance(author_maintenance, dict):
        model_author = first_non_empty(str(author_maintenance.get("bio", "")))
    generated = first_non_empty(model_author, section_value(result, "author"))
    if generated:
        return short_text(generated, 1000), "AI已补全"
    return "　", "待补全失败"


def build_author_maintenance(facts: dict, result: dict, editor_name: str = "叶文涛") -> dict:
    bio, bio_status = build_author_bio(facts, result)
    name = extract_author_name_from_facts(facts)
    work_unit = author_field(facts, "工作单位", "单位名称", "作者单位", "所在单位")
    major = author_field(facts, "专业", "作者专业", "适用专业")
    writing_direction = author_field(facts, "著作方向", "研究方向", "专业方向")
    return {
        "enabled": bool(name),
        "type": "个人",
        "name": name,
        "contactor": editor_name,
        "gender": author_field(facts, "性别"),
        "certificateType": author_field(facts, "证件类型") or "其他",
        "certificateNo": blank_if_missing(author_field(facts, "证件号", "身份证号")),
        "title": blank_if_missing(author_field(facts, "职务/职称", "职称", "职务")),
        "education": blank_if_missing(author_field(facts, "学历", "学位")),
        "major": blank_if_missing(major),
        "graduateSchool": blank_if_missing(author_field(facts, "毕业院校", "毕业学校")),
        "workUnit": blank_if_missing(work_unit),
        "unitAddress": blank_if_missing(author_field(facts, "单位地址")),
        "unitZip": blank_if_missing(author_field(facts, "单位邮编", "邮编", "邮政编码")),
        "unitContactor": blank_if_missing(author_field(facts, "单位联系人")),
        "unitFax": blank_if_missing(author_field(facts, "单位传真")),
        "phone1": blank_if_missing(author_field(facts, "联系电话", "联系电话1", "手机", "电话")),
        "phone2": blank_if_missing(author_field(facts, "联系电话2")),
        "fax": blank_if_missing(author_field(facts, "传真")),
        "email": blank_if_missing(author_field(facts, "电子邮箱", "邮箱", "电子邮件")),
        "postalCode": blank_if_missing(author_field(facts, "邮编", "邮政编码")),
        "address": blank_if_missing(author_field(facts, "通信地址", "通讯地址", "联系地址")),
        "workResume": blank_if_missing(author_field(facts, "工作简历或单位简介", "工作简历", "个人简历", "个人简介")),
        "academicOrganizations": blank_if_missing(author_field(facts, "参加的学术组织及任职情况", "学术或教育组织任职")),
        "researchProjects": blank_if_missing(author_field(facts, "科研或教研项目经历", "项目经历", "教研项目")),
        "awards": blank_if_missing(author_field(facts, "科研或教学工作及获奖情况", "获奖情况", "教学获奖")),
        "writingDirection": blank_if_missing(writing_direction or major),
        "publications": blank_if_missing(author_field(facts, "主要著作出版情况", "已出版教材或著作", "著作出版情况")),
        "bio": bio,
        "bioStatus": bio_status,
    }


def substantive_author_value(value) -> bool:
    text = strip_indent(str(value or ""))
    return bool(text and text not in {"[已隐藏]", "[REDACTED]"})


def merge_private_author_maintenance(current: dict, extracted: dict) -> dict:
    merged = dict(current or {})
    for key, value in (extracted or {}).items():
        if key in {"bio", "bioStatus"} and substantive_author_value(merged.get(key)):
            continue
        if substantive_author_value(value) and not substantive_author_value(merged.get(key)):
            merged[key] = value
    return merged


def hydrate_author_payload_from_project_application(
    payload: dict, user_id: int, project_id: str
) -> dict:
    application_files = [
        item
        for item in APP_STORE.list_project_files(user_id, project_id)
        if item.get("kind") == "application"
    ]
    if not application_files:
        return payload

    application_path = PROJECT_FILE_STORE.resolve(application_files[0]["storagePath"])
    facts = load_extractor().build_payload(application_path, include_sensitive=True)
    hydrated = deepcopy(payload)
    topic = merge_bpm_topic(hydrated)
    current = topic.get("authorMaintenance") if isinstance(topic.get("authorMaintenance"), dict) else {}
    extracted = build_author_maintenance(
        facts,
        {"authorMaintenance": current},
        editor_name_from_payload(hydrated),
    )
    author = merge_private_author_maintenance(current, extracted)
    author["enabled"] = True

    gender = strip_indent(str(author.get("gender", "")))
    email = strip_indent(str(author.get("email", "")))
    if gender not in {"男", "女"}:
        raise ValueError("原始选题申报表缺少有效作者性别，不能新增 BPM 作译者。")
    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        raise ValueError("原始选题申报表缺少合法电子邮箱，不能新增 BPM 作译者。")

    topic["authorMaintenance"] = author
    topic["authorName"] = first_non_empty(topic.get("authorName", ""), author.get("name", ""))
    hydrated["bpmTopic"] = topic
    hydrated["authorMaintenance"] = author
    return hydrated


def report_scores(result_or_payload: dict) -> dict:
    names = {
        "选题内容": "scoreContent",
        "作者情况": "scoreAuthor",
        "策划过程与可行性": "scoreFeasibility",
        "获奖潜质": "scoreAward",
        "成本与盈利估算": "scoreProfit",
        "市场定位与营销": "scoreMarketing",
    }
    scores = {}
    total = 0
    for item in result_or_payload.get("scores", []):
        if not isinstance(item, list) or len(item) < 3:
            continue
        key = names.get(item[0])
        if not key:
            continue
        try:
            value = int(item[2])
        except Exception:
            value = 0
        scores[key] = value
        total += value
    if scores:
        scores["scoreTotal"] = total
    return scores


def build_bpm_topic(facts: dict, result: dict) -> dict:
    today = date.today()
    content = section_value(result, "content")
    author = section_value(result, "author")
    marketing = section_value(result, "marketing")
    title = first_non_empty(
        field_value(facts, "教材名称", "选题名称", "选题名", "书名"),
        result.get("title", ""),
        "选题策划报告",
    )
    bpm_fields = normalize_bpm_fields(result.get("bpmFields") or result.get("bpm_fields"))
    brief = first_non_empty(bpm_fields.get("brief", ""), field_value(facts, "内容简介"), content)
    reader = first_non_empty(
        bpm_fields.get("reader", ""),
        field_value(facts, "读者 对象", "读者对象", "读者定位"),
        marketing,
    )
    feature = first_non_empty(bpm_fields.get("feature", ""), field_value(facts, "本教材特色和优势", "特色和优势"), content)
    compare = first_non_empty(bpm_fields.get("compare", ""), field_value(facts, "国内外同类教材比较", "同类教材比较"), marketing)
    author_name = extract_author_name_from_facts(facts)
    book_name = safe_name(title).replace("-选题策划报告", "")
    detailed_classification = infer_bpm_detailed_classification(title, facts, content, marketing)
    model_classification = result.get("bpmClassification") or result.get("bpm_classification") or {}
    class3 = first_non_empty(str(model_classification.get("class3", "")), detailed_classification.get("class3", ""))
    class4 = first_non_empty(str(model_classification.get("class4", "")), detailed_classification.get("class4", ""))

    editor_name = editor_name_from_payload({})
    return {
        **fixed_bpm_defaults(),
        **detailed_classification,
        **BPM_FIXED_CLASSIFICATION,
        "class3": class3,
        "class4": class4,
        "bookName": book_name,
        "authorName": author_name,
        "authorCode": "",
        "authorId": "",
        "brief": short_text(brief),
        "reader": short_text(reader),
        "feature": short_text(feature),
        "compare": short_text(compare),
        "compareSources": bpm_fields.get("compareSources", []),
        "scriptDate": normalize_bpm_date(field_value(facts, "预定交稿时间", "交稿日期"), today + timedelta(days=90)),
        "makingDate": normalize_bpm_date(field_value(facts, "预计发稿日期"), today + timedelta(days=150)),
        "publishDate": normalize_bpm_date(field_value(facts, "预计出版日期", "出版日期"), today + timedelta(days=270)),
        "authorMaintenance": build_author_maintenance(facts, result, editor_name),
        **report_scores(result),
    }


def merge_bpm_topic(payload: dict) -> dict:
    topic = dict(payload.get("bpmTopic") or {})
    bpm_fields = normalize_bpm_fields(payload.get("bpmFields") or payload.get("bpm_fields"))
    editor_name = editor_name_from_payload(payload)
    topic.update(BPM_FIXED_CLASSIFICATION)
    topic["class3"] = first_non_empty(str(topic.get("class3", "")), "020101")
    topic["class4"] = first_non_empty(str(topic.get("class4", "")), "02010103")
    topic["bookName"] = safe_name(payload.get("title") or topic.get("bookName") or "选题策划报告")
    author_name_source = str(topic.get("authorName", ""))
    if isinstance(topic.get("authorMaintenance"), dict):
        author_name_source = first_non_empty(author_name_source, topic["authorMaintenance"].get("name", ""))
    topic["authorName"] = clean_person_name(author_name_source)
    topic["projectEditor"] = editor_name
    topic["editor"] = editor_name
    for identity_key in ("projectEditorNo", "projectEditorUid", "editorNo", "editorUid"):
        topic.pop(identity_key, None)
    sections = {section.get("key"): strip_indent(section.get("text", "")) for section in payload.get("sections", [])}
    topic["brief"] = short_text(
        first_non_empty(bpm_fields.get("brief", ""), topic.get("brief", ""), sections.get("content", ""))
    )
    topic["reader"] = short_text(
        first_non_empty(bpm_fields.get("reader", ""), topic.get("reader", ""), sections.get("marketing", ""))
    )
    topic["feature"] = short_text(first_non_empty(bpm_fields.get("feature", ""), topic.get("feature", ""), sections.get("content", "")))
    topic["compare"] = short_text(first_non_empty(bpm_fields.get("compare", ""), topic.get("compare", ""), sections.get("marketing", "")))
    topic["compareSources"] = bpm_fields.get("compareSources") or normalize_compare_sources(
        topic.get("compareSources")
    )
    if isinstance(topic.get("authorMaintenance"), dict):
        topic["authorMaintenance"]["contactor"] = editor_name
        topic["authorMaintenance"].pop("contactorUid", None)
        topic["authorMaintenance"].pop("contactorDeptId", None)
    topic.update(report_scores(payload))
    return topic


class BpmRunError(RuntimeError):
    def __init__(self, message: str, *, result: dict | None = None):
        super().__init__(message)
        self.result = result if isinstance(result, dict) else None


def run_bpm_script(payload: dict, mode: str, failure_label: str) -> dict:
    if not BPM_SCRIPT.exists():
        raise FileNotFoundError(f"BPM skill 脚本不存在: {BPM_SCRIPT}")
    credentials = payload.get("bpm") or {}
    bpm_user = (credentials.get("user") or os.environ.get("BPM_USER") or "").strip()
    bpm_password = (credentials.get("password") or os.environ.get("BPM_PASSWORD") or "").strip()
    bpm_url = (credentials.get("url") or os.environ.get("BPM_URL") or "http://bpm.phei.com.cn:8088/portal/r/w").strip()
    if not bpm_user or not bpm_password:
        raise ValueError("请填写 BPM 账号和密码。")

    topic = merge_bpm_topic(payload)
    run_root = ROOT / "output" / "bpm-runs"
    run_root.mkdir(parents=True, exist_ok=True)
    run_dir = run_root / f"{datetime.now().strftime('%Y%m%d-%H%M%S')}-{uuid.uuid4().hex[:8]}"
    run_dir.mkdir(parents=True, exist_ok=True)
    topic_path = run_dir / "topic.json"
    topic_path.write_text(json.dumps(topic, ensure_ascii=False, indent=2), encoding="utf-8")
    env = os.environ.copy()
    env.update(
        {
            "BPM_USER": bpm_user,
            "BPM_PASSWORD": bpm_password,
            "BPM_URL": bpm_url,
            "NODE_PATH": str(NODE_MODULES),
        }
    )
    completed = subprocess.run(
        ["node", str(BPM_SCRIPT), mode, str(topic_path)],
        cwd=str(ROOT),
        env=env,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        timeout=180,
        check=False,
    )
    (run_dir / "stdout.txt").write_text(completed.stdout or "", encoding="utf-8")
    (run_dir / "stderr.txt").write_text(completed.stderr or "", encoding="utf-8")
    result = None
    if completed.stdout.strip():
        try:
            parsed = json.loads(completed.stdout)
            if isinstance(parsed, dict):
                result = parsed
                result["runDir"] = str(run_dir)
        except json.JSONDecodeError:
            result = None
    if completed.returncode != 0:
        detail = (
            str((result or {}).get("error") or "").strip()
            or completed.stderr.strip()
            or completed.stdout.strip()
            or f"{failure_label}失败"
        )
        raise BpmRunError(
            f"{detail}\n运行记录已保存：{run_dir}", result=result
        )
    if result is None:
        raise BpmRunError(
            f"{failure_label}返回结果无法解析。运行记录已保存：{run_dir}"
        )
    if not result.get("ok"):
        raise BpmRunError(
            f"{failure_label}未验证成功。运行记录已保存：{run_dir}",
            result=result,
        )
    return result


def run_bpm_topic_submit(payload: dict) -> dict:
    result = run_bpm_script(payload, "submit-topic", "BPM 选题暂存")
    if "worklist_verified" not in result.get("milestones", []):
        raise BpmRunError(
            "BPM 选题暂存未完成工作列表验证（缺少 worklist_verified）。",
            result=result,
        )
    if not result.get("title"):
        expected = result.get("expectedBookName") or merge_bpm_topic(payload).get("bookName") or "选题"
        raise BpmRunError(
            f"BPM 选题暂存未验证成功：没有在待办列表找到《{expected}》。",
            result=result,
        )
    return result


def run_bpm_author_submit(payload: dict) -> dict:
    topic = merge_bpm_topic(payload)
    author = topic.get("authorMaintenance") if isinstance(topic.get("authorMaintenance"), dict) else {}
    author_name = clean_person_name(str(author.get("name", "")))
    author_bio = str(author.get("bio", "")).strip()
    if not author_name:
        raise ValueError("没有识别到有效作者姓名，请检查申报表中的作者姓名字段。")
    if not author_bio:
        raise ValueError("作者简介为空，不能新增作译者。")
    author["enabled"] = True
    topic["authorMaintenance"] = author
    author_payload = deepcopy(payload)
    author_payload["bpmTopic"] = topic
    result = run_bpm_script(author_payload, "submit-author", "BPM 作译者保存")
    if not result.get("authorCode") or result.get("verified") is not True:
        raise BpmRunError(
            "BPM 作译者保存缺少作译者编码或列表验证，不能标记为成功。",
            result=result,
        )
    return result


def run_bpm_submit(payload: dict) -> dict:
    return run_bpm_topic_submit(payload)


def now_iso() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def bpm_topic_preview(topic: dict) -> dict:
    keys = [
        "bookName",
        "authorName",
        "class1",
        "class2",
        "class3",
        "class4",
        "gbClass",
        "readLevel",
        "bwClass",
        "bwCipClass",
        "readerNum",
        "scriptSource",
        "words",
        "price",
        "remPayMode",
        "remStandard",
        "digital",
        "eRemPayMode",
        "totalNum",
        "firstNum",
        "projectEditor",
        "editor",
        "feature",
        "compare",
        "scoreTotal",
    ]
    return {key: topic.get(key, "") for key in keys}


def redact_bpm_secret(value, password: str):
    return APP_STORE.sanitize_job_data(value, known_secrets=(password,))


def trusted_bpm_payload(payload: dict, user: dict) -> tuple[dict, dict]:
    credentials = APP_STORE.get_integration_credentials(user["id"], "phei_bpm")
    if not credentials:
        raise ValueError("请先保存 BPM 账号和密码。")
    trusted_payload = deepcopy(payload)
    trusted_payload["editorName"] = user["display_name"]
    trusted_payload["bpm"] = {
        "url": os.environ.get("BPM_URL", "http://bpm.phei.com.cn:8088/portal/r/w"),
        "user": credentials["account"],
        "name": user["display_name"],
        "password": credentials["password"],
    }
    return trusted_payload, credentials


def business_bpm_payload(payload: dict) -> dict:
    if not isinstance(payload, dict):
        raise ValueError("BPM 请求体必须是对象。")
    sensitive_keys = {
        "bpm",
        "credentials",
        "bpmcredentials",
        "password",
        "bpmpassword",
        "bpm_password",
    }

    def remove_sensitive(value):
        if isinstance(value, list):
            return [remove_sensitive(item) for item in value]
        if not isinstance(value, dict):
            return value
        return {
            key: remove_sensitive(item)
            for key, item in value.items()
            if str(key).replace("-", "").replace("_", "").lower() not in sensitive_keys
        }

    business_payload = remove_sensitive(deepcopy(payload))
    business_payload.pop("editorName", None)
    business_payload.pop("createdBy", None)
    return business_payload


def project_report_payload(project: dict, user: dict) -> dict:
    summary = summarize_project_state(project.get("state") or {}, user["display_name"])
    state = summary["state"]
    scores = state["scoreItems"]
    total = 0
    for item in scores:
        try:
            total += float(item[2])
        except (TypeError, ValueError, IndexError):
            pass
    if float(total).is_integer():
        total = int(total)
    return {
        "projectId": project["id"],
        "title": summary["title"],
        "editorName": user["display_name"],
        "createdBy": user["display_name"],
        "sections": state["sections"],
        "scores": scores,
        "scoreItems": scores,
        "total": total,
        "bpmTopic": state["bpmTopic"],
        "authorMaintenance": state["authorMaintenance"],
    }


def load_project_job_payload(user: dict, project_id: str, job_type: str) -> tuple[dict, dict]:
    if not isinstance(project_id, str) or not project_id.strip():
        raise ValueError("请先选择要处理的选题项目。")
    project = APP_STORE.get_project(user["id"], project_id.strip())
    if project is None or project.get("archivedAt"):
        raise ValueError("未找到当前用户的选题项目。")
    payload = project_report_payload(project, user)
    if job_type == "topic":
        preflight = build_bpm_preflight(
            payload_to_project_state(payload),
            bpm_configured=APP_STORE.has_integration_credentials(user["id"], "phei_bpm"),
        )
        if not preflight["ready"]:
            blockers = [
                row["label"]
                for row in preflight["rows"]
                if row["status"] == "blocking_missing"
            ]
            raise ValueError("项目尚未通过 BPM 预检：" + "、".join(blockers))
    return payload, project


def payload_to_project_state(payload: dict) -> dict:
    return normalize_project_state(
        {
            "sections": payload.get("sections"),
            "scoreItems": payload.get("scoreItems", payload.get("scores")),
            "bpmTopic": payload.get("bpmTopic"),
            "authorMaintenance": payload.get("authorMaintenance"),
        }
    )


def update_project_job_status(
    user_id: int, project_id: str | None, job_type: str, status: str
) -> None:
    if not project_id:
        return
    status_key = "author_status" if job_type == "author" else "bpm_status"
    for _ in range(3):
        project = APP_STORE.get_project(user_id, project_id)
        if project is None:
            return
        try:
            APP_STORE.update_project(
                user_id,
                project_id,
                project["version"],
                {status_key: status},
            )
            return
        except ProjectVersionConflict:
            continue
        except Exception as error:
            runtime_log(
                f"project status update failure project={project_id} status={status} "
                f"type={type(error).__name__}"
            )
            return
    runtime_log(f"project status update conflict project={project_id} status={status}")


def create_bpm_job(payload: dict, user: dict, job_type: str = "topic") -> dict:
    if job_type not in {"topic", "author"}:
        raise ValueError(f"不支持的 BPM 任务类型：{job_type}")
    requested_project_id = payload.get("projectId") if isinstance(payload, dict) else None
    project = None
    if requested_project_id:
        trusted_project_payload, project = load_project_job_payload(
            user, requested_project_id, job_type
        )
        business_payload = business_bpm_payload(trusted_project_payload)
    else:
        business_payload = business_bpm_payload(payload)
    if not APP_STORE.has_integration_credentials(user["id"], "phei_bpm"):
        raise ValueError("请先保存 BPM 账号和密码。")
    validation_payload = deepcopy(business_payload)
    validation_payload["editorName"] = user["display_name"]
    topic = merge_bpm_topic(validation_payload)
    if job_type == "topic":
        sections = validation_payload.get("sections") or []
        if len(sections) < 6 or any(not section.get("text") for section in sections):
            raise ValueError("请先生成完整的一到六部分报告内容。")
        unconfirmed = [section for section in sections if not section.get("confirmed")]
        if unconfirmed:
            raise ValueError(f"还有 {len(unconfirmed)} 段未确认，不能加入 BPM 队列。")
        job_title = topic.get("bookName") or validation_payload.get("title") or "选题策划报告"
        job_label = "选题填报"
    else:
        author = topic.get("authorMaintenance") if isinstance(topic.get("authorMaintenance"), dict) else {}
        author_name = clean_person_name(str(author.get("name", "")))
        if not author_name:
            raise ValueError("没有识别到有效作者姓名，请检查申报表中的作者姓名字段。")
        if not str(author.get("bio", "")).strip():
            raise ValueError("作者简介为空，不能新增作译者。")
        job_title = author_name
        job_label = "作译者维护"
    user_id = user["id"]
    job = APP_STORE.create_job(
        user_id,
        job_type,
        job_title,
        {
            "createdBy": user["display_name"],
            "topicPreview": bpm_topic_preview(topic),
        },
        project_id=None if project is None else project["id"],
    )
    job_id = job["id"]
    APP_STORE.append_job_log(job_id, user_id, f"已加入{job_label}队列")
    if project is not None:
        business_payload["projectId"] = project["id"]
        update_project_job_status(user_id, project["id"], job_type, "queued")
    JOB_QUEUE.put((job_id, user_id, job_type, business_payload))
    return APP_STORE.list_jobs(user_id, limit=1)[0]


def log_job_storage_failure(job_id: str, operation: str, error: Exception):
    runtime_log(
        f"bpm-job storage failure job={job_id} operation={operation} type={type(error).__name__}"
    )


def update_bpm_job(job_id: str, user_id: int, status: str, result=None, error_summary=None):
    try:
        APP_STORE.update_job(job_id, user_id, status, result=result, error_summary=error_summary)
    except Exception as error:
        log_job_storage_failure(job_id, "update", error)


def append_bpm_job_log(job_id: str, user_id: int, message: str):
    try:
        APP_STORE.append_job_log(job_id, user_id, message)
    except Exception as error:
        log_job_storage_failure(job_id, "append-log", error)


def process_bpm_job(job_id: str, user_id: int, job_type: str, payload: dict):
    job_label = "作译者维护" if job_type == "author" else "选题填报"
    password = ""
    trusted_payload = {}
    credentials = {}
    update_bpm_job(job_id, user_id, "running")
    project_id = payload.get("projectId") if isinstance(payload, dict) else None
    update_project_job_status(user_id, project_id, job_type, "running")
    try:
        user = APP_STORE.get_user_by_id(user_id)
        if not user:
            raise ValueError("当前用户不存在，不能执行 BPM 任务。")
        trusted_payload, credentials = trusted_bpm_payload(payload, user)
        if job_type == "author" and project_id:
            trusted_payload = hydrate_author_payload_from_project_application(
                trusted_payload, user_id, project_id
            )
        password = credentials.get("password", "")
        append_bpm_job_log(job_id, user_id, f"开始登录 BPM 并执行{job_label}")
        result = (
            run_bpm_author_submit(trusted_payload)
            if job_type == "author"
            else run_bpm_topic_submit(trusted_payload)
        )
        if job_type == "topic" and "worklist_verified" not in result.get("milestones", []):
            raise BpmRunError(
                "BPM 选题填报缺少 worklist_verified 里程碑，不能标记为成功。",
                result=result,
            )
        public_result = redact_bpm_secret(result, password)
        update_bpm_job(job_id, user_id, "succeeded", result=public_result)
        update_project_job_status(user_id, project_id, job_type, "succeeded")
        title = result.get("authorName") or result.get("title") or result.get("expectedBookName") or job_label
        append_bpm_job_log(job_id, user_id, f"{job_label}完成：{redact_bpm_secret(str(title), password)}")
    except Exception as error:
        error_summary = redact_bpm_secret(str(error), password)
        failure_result = getattr(error, "result", None)
        public_failure_result = redact_bpm_secret(failure_result, password)
        update_bpm_job(
            job_id,
            user_id,
            "failed",
            result=public_failure_result,
            error_summary=error_summary,
        )
        update_project_job_status(user_id, project_id, job_type, "failed")
        append_bpm_job_log(job_id, user_id, f"{job_label}失败：{error_summary}")
    finally:
        clear_bpm_job_password(trusted_payload)
        if isinstance(credentials, dict):
            credentials.clear()


def clear_bpm_job_password(payload):
    try:
        if not isinstance(payload, dict):
            return
        bpm = payload.get("bpm")
        if isinstance(bpm, dict):
            bpm.pop("password", None)
    except Exception:
        pass


def bpm_job_worker(work_queue=None, stop_event=None):
    work_queue = JOB_QUEUE if work_queue is None else work_queue
    while True:
        if stop_event is not None and stop_event.is_set():
            return
        try:
            item = work_queue.get(timeout=0.1 if stop_event is not None else None)
        except queue.Empty:
            continue
        payload = {}
        job_id = "unknown"
        try:
            job_id, user_id, job_type, payload = item
            process_bpm_job(job_id, user_id, job_type, payload)
        except Exception as error:
            runtime_log(f"bpm-job worker failure job={job_id} type={type(error).__name__}")
        finally:
            try:
                clear_bpm_job_password(payload)
            finally:
                work_queue.task_done()


threading.Thread(target=bpm_job_worker, daemon=True).start()


def extract_json_from_text(text: str) -> dict:
    text = text.strip()
    fence = re.search(r"```(?:json)?\s*(\{.*?\})\s*```", text, flags=re.S)
    if fence:
        text = fence.group(1)
    else:
        start = text.find("{")
        end = text.rfind("}")
        if start >= 0 and end > start:
            text = text[start : end + 1]
    return json.loads(text)


def normalize_model_url(url: str) -> str:
    url = (url or DEFAULT_MODEL_URL).strip().rstrip("/")
    if not url.endswith("/chat/completions"):
        if url.endswith("/v1"):
            url += "/chat/completions"
        else:
            url += "/v1/chat/completions"
    return url


def build_prompt(facts: dict, book_context: list[dict] | None = None) -> str:
    report_rules = read_text(REPORT_RULES)
    scoring_rules = read_text(SCORING_RULES)
    compact = compact_facts(facts)
    if book_context is None:
        book_context = book_comparison_search_context(facts)
    if book_context:
        compact["book_comparison_search_context"] = book_context
    source = json.dumps(compact, ensure_ascii=False, indent=2)
    return f"""
你是电子工业出版社教育出版板块的资深策划编辑。请严格依据申报表事实和下方 skill 规则，生成供出版社领导审议的选题策划报告一到六部分、建议评分和待确认信息。报告用于论证选题的出版价值和实施可行性，帮助领导判断是否批准立项。

	硬性要求：
	1. 只输出 JSON，不要输出 Markdown 解释。
	2. 一到六每个正文段落必须以两个全角空格开头：`　　`。
	3. 不得编造定价、版税、包销、资助、毛利润、获奖、版权授权等精确信息；缺失则写入 pending_questions。
	4. 第五部分“成本与盈利估算”必须使用下方固定文本，不得改写、补充或替换数字：
	{FIXED_PROFIT_SECTION_TEXT}
	5. 除第五部分外，一、二、三、四、六必须先依据申报表生成内容，再按下方“范文润色规则”重写成更成熟的领导审议材料；不要保留流水账、表格腔或申报表原话堆砌。
	6. 一到六正文只写来源能够支持的事实与合理判断。某项信息缺失时，直接略去，不得在正文中写“申报表未提供”“申报表中没有体现”“目前尚不明确”“缺少相关依据”“有待进一步确认”等缺项说明。
	7. 一到六正文用于论证“为什么值得出版、现有基础为什么能够支撑实施”，不得向责任编辑、策划编辑或作者布置工作，不得出现“建议责任编辑”“建议在组稿阶段”“建议后续跟进”“建议进一步确认”“不建议在论证中强调”“待作者补充后再”等内部工作指令。
	8. BPM 填报分类中，一级分类固定为“教育”，二级分类固定为“本科研究生”，国标分类固定为“G”，层次固定为“高等理工”；你只需要根据选题内容理解，给出最合适的三级分类和四级分类建议，可填 BPM 下拉框文本或代码，不确定则留空。
	9. 额外生成四个 BPM 填报专用字段，它们不是策划报告正文摘要，不得直接复制第一部分或第六部分：
	   - `bpmFields.brief` 是“内容简介”，用 250-400 字说明本书讲什么、按什么逻辑组织、读者能够获得什么能力。只写书稿内容，不写市场宣传、编辑建议或缺项说明。
	   - `bpmFields.reader` 是“读者对象”，只写一句“本书适合大学XXX专业本科生、研究生，以及从事XXX相关工作的人员使用。”根据选题选择专业和相关工作，不加入高职学生、继续教育学员等额外人群。
	   - `bpmFields.feature` 是“选题特色”，用 250-400 字说明内容体系、教材组织、实践教学和行业前沿价值，回答为什么本书值得出版；不写“建议责任编辑”“后续补充”等内部工作意见。
	   - `bpmFields.compare` 是“同类选题比较”。只能从 `book_comparison_search_context` 中恰好选择 2 本可核验的真实图书，逐本按“书名（作者，出版社，年份）—优点—相较本选题的覆盖差异”写，末尾总结本选题的差异化定位。所谓不足只能写成基于公开简介和目录的相对覆盖差异，不得评价写作质量，不得编造销量、排名或市场数据。
	   - `bpmFields.compareSources` 恰好保存上述 2 本书的来源信息，至少包含 `title` 和 `url`；书名及网址必须来自 `book_comparison_search_context`，不得自行生成网址。来源网址只供后台追溯，不要写进 `compare` 正文。
	   - 如果联网检索得到的可核验图书不足 2 本，`compare` 和 `compareSources` 留空，不得用模型记忆虚构书目。
	10. 额外生成 `authorMaintenance.bio`：如果申报表已有作者简介则提炼为 50-1000 字；如果没有作者简介但 `author_official_search_context` 中有学校/学院/单位官网摘要，则优先依据这些公开摘要和申报表事实写作者简介；如果两者都不足，则只根据申报表中作者单位、职称、学历、研究/教学经历、项目、获奖、著作等已知事实，写一段可用于 BPM 作译者维护的作者简介。不得编造精确头衔、项目名称、获奖名称或联系方式。作者姓名不要从“作者情况”正文推断；如申报表没有独立姓名字段但有“合作者情况简介（姓名、年龄、职称、工作单位等）”，以该栏首位作者姓名为准。
	11. 评分按保守口径，总分通常控制在 60-70；除非材料非常强，不要超过 70。
	12. 输出字段必须符合下面 JSON 结构：
{{
  "title": "选题名称",
  "sections": [
    {{"key": "content", "title": "一、选题内容", "text": "..."}},
    {{"key": "author", "title": "二、作者情况", "text": "..."}},
    {{"key": "feasibility", "title": "三、策划过程与可行性", "text": "..."}},
    {{"key": "award", "title": "四、获奖潜质", "text": "..."}},
    {{"key": "profit", "title": "五、成本与盈利估算", "text": "..."}},
    {{"key": "marketing", "title": "六、市场定位与营销", "text": "..."}}
  ],
  "scores": [
    ["选题内容", 35, 30],
    ["作者情况", 10, 5],
    ["策划过程与可行性", 5, 5],
    ["获奖潜质", 5, 0],
    ["成本与盈利估算", 35, 17],
    ["市场定位与营销", 10, 8]
  ],
  "pending_questions": ["..."],
	  "bpmClassification": {{"class3": "三级分类建议", "class4": "四级分类建议"}},
	  "bpmFields": {{
	    "brief": "250-400字内容简介",
	    "reader": "本书适合大学XXX专业本科生、研究生，以及从事XXX相关工作的人员使用。",
	    "feature": "250-400字选题特色",
	    "compare": "两本真实图书的逐本比较及本选题差异化定位",
	    "compareSources": [
	      {{"title": "第一本书名", "url": "来源网址", "author": "作者", "publisher": "出版社", "year": "年份"}},
	      {{"title": "第二本书名", "url": "来源网址", "author": "作者", "publisher": "出版社", "year": "年份"}}
	    ]
	  }},
  "authorMaintenance": {{
    "bio": "作者简介，50-1000字"
  }}
}}

	【报告生成规则】
	{report_rules}

	【范文润色规则】
	{AI_REPORT_STYLE_GUIDE}

	【评分规则】
{scoring_rules}

【申报表抽取信息】
{source}
""".strip()


def call_model(prompt: str, model_url: str, api_key: str, model: str) -> dict:
    endpoint = normalize_model_url(model_url)
    resolved_api_key = api_key.strip()
    if not resolved_api_key:
        raise ModelConfigurationError(MODEL_CONFIGURATION_ERROR)
    payload = {
        "model": (model or DEFAULT_MODEL).strip(),
        "messages": [
            {"role": "system", "content": "你只输出严格 JSON。"},
            {"role": "user", "content": prompt},
        ],
        "temperature": 0.2,
        "stream": False,
    }
    data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    request = urllib.request.Request(
        endpoint,
        data=data,
        headers={
            "Content-Type": "application/json",
            "Authorization": f"Bearer {resolved_api_key}",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=120) as response:
            body = response.read().decode("utf-8")
    except urllib.error.HTTPError as exc:
        runtime_log(f"model request failed status={exc.code}")
        raise ModelServiceError(MODEL_REQUEST_ERROR) from exc
    except urllib.error.URLError as exc:
        runtime_log(f"model request failed type={type(exc.reason).__name__}")
        raise ModelServiceError(MODEL_REQUEST_ERROR) from exc
    try:
        result = json.loads(body)
        content = result["choices"][0]["message"]["content"]
        return extract_json_from_text(content)
    except (KeyError, TypeError, json.JSONDecodeError) as exc:
        runtime_log(f"model response invalid type={type(exc).__name__}")
        raise ModelServiceError(MODEL_REQUEST_ERROR) from exc


def call_complete_report_model(prompt: str) -> dict:
    return call_model(prompt, DEFAULT_MODEL_URL, DEFAULT_API_KEY, DEFAULT_MODEL)


def call_fast_model(prompt: str) -> dict:
    return call_model(prompt, DEFAULT_MODEL_URL, DEFAULT_API_KEY, FAST_MODEL)


def find_report_tone_violations(result: dict) -> list[dict]:
    violations = []
    sections = result.get("sections") if isinstance(result, dict) else []
    for section in sections or []:
        if not isinstance(section, dict) or section.get("key") == "profit":
            continue
        text = str(section.get("text", "") or "")
        matches = []
        for pattern in REPORT_TONE_BANNED_PATTERNS:
            matches.extend(match.group(0) for match in re.finditer(pattern, text))
        matches = list(dict.fromkeys(matches))
        if matches:
            violations.append(
                {
                    "key": str(section.get("key", "") or ""),
                    "title": str(section.get("title", "") or ""),
                    "matches": matches,
                }
            )
    return violations


def build_report_tone_rewrite_prompt(facts: dict, result: dict, violations: list[dict]) -> str:
    violation_keys = {item["key"] for item in violations}
    sections = [
        section
        for section in result.get("sections", [])
        if isinstance(section, dict) and section.get("key") in violation_keys
    ]
    source = json.dumps(compact_facts(facts), ensure_ascii=False, indent=2)
    section_source = json.dumps(sections, ensure_ascii=False, indent=2)
    matched = json.dumps(violations, ensure_ascii=False, indent=2)
    return f"""
你是电子工业出版社教育出版板块的资深策划编辑。下面部分选题策划报告包含不适合提交领导审议的内部工作指令或缺项说明，请仅重写这些部分。

重写要求：
1. 报告用于论证选题的出版价值和实施可行性，帮助领导判断是否批准立项。
2. 保留来源能够支持的事实，将内容写成积极、稳健、有依据的出版价值判断。
3. 不得增加来源中没有的作者经历、奖项、市场数据、版权合作或资源成果。
4. 某项信息缺失时直接略去，不得说明“申报表未提供”“申报表中没有体现”“缺少相关依据”或“有待进一步确认”。
5. 不得向责任编辑、策划编辑或作者布置工作，不得出现“建议责任编辑”“建议在组稿阶段”“建议后续跟进”“不建议强调”或“待作者补充”等指令。
6. 每个正文段落以两个全角空格开头。只输出严格 JSON，不要解释修改过程。
7. 仅输出需要重写的 sections，结构如下：
{{
  "sections": [
    {{"key": "原key", "title": "原标题", "text": "重写后的正文"}}
  ]
}}

【检测到的问题】
{matched}

【需要重写的部分】
{section_source}

【申报表抽取信息】
{source}
""".strip()


def merge_rewritten_report_sections(result: dict, rewritten: dict, violation_keys: set[str]) -> dict:
    replacements = {
        str(section.get("key", "")): section
        for section in rewritten.get("sections", [])
        if isinstance(section, dict)
        and section.get("key") in violation_keys
        and str(section.get("text", "") or "").strip()
    }
    merged = deepcopy(result)
    for section in merged.get("sections", []):
        replacement = replacements.get(str(section.get("key", "")))
        if replacement:
            section["title"] = replacement.get("title") or section.get("title")
            section["text"] = replacement["text"]
    return merged


def revise_report_tone_if_needed(facts: dict, result: dict) -> dict:
    revised = deepcopy(result)
    for attempt in range(2):
        violations = find_report_tone_violations(revised)
        if not violations:
            return revised
        runtime_log(
            "report tone rewrite start "
            f"attempt={attempt + 1} sections={','.join(item['key'] for item in violations)}"
        )
        rewritten = call_complete_report_model(
            build_report_tone_rewrite_prompt(facts, revised, violations)
        )
        revised = merge_rewritten_report_sections(
            revised,
            rewritten,
            {item["key"] for item in violations},
        )
    remaining = find_report_tone_violations(revised)
    if remaining:
        runtime_log(
            "report tone rewrite failed "
            f"sections={','.join(item['key'] for item in remaining)}"
        )
        raise ModelServiceError("报告正文仍包含不适合领导审议的表述，请重新生成。")
    return revised


CORE_FIELD_ALIASES = {
    "选题名称": ("教材名称", "选题名称", "选题名", "书名"),
    "姓名": ("作者姓名", "姓名", "主要作（译）者姓名", "第一作者姓名", "主编"),
}

RECOVERY_MODEL_KEYS = {
    "选题名称": "bookName",
    "姓名": "authorName",
}

SENSITIVE_SOURCE_LABELS = (
    "身份证",
    "证件号",
    "联系电话",
    "电话",
    "手机",
    "电子邮箱",
    "邮箱",
    "email",
    "邮政编码",
    "邮编",
    "通信地址",
    "通讯地址",
)


def missing_core_fields(facts: dict) -> list[str]:
    missing = []
    if not field_value(facts, *CORE_FIELD_ALIASES["选题名称"]):
        missing.append("选题名称")
    if not extract_author_name_from_facts(facts):
        missing.append("姓名")
    return missing


def is_sensitive_source_label(value: str) -> bool:
    compact = re.sub(r"\s+", "", str(value or "")).lower()
    return any(label.lower() in compact for label in SENSITIVE_SOURCE_LABELS)


def scrub_private_source_value(value: str) -> str:
    text = str(value or "")
    text = re.sub(r"(?i)\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", "[已隐藏]", text)
    text = re.sub(r"(?<!\d)1[3-9]\d{9}(?!\d)", "[已隐藏]", text)
    text = re.sub(r"(?<!\d)\d{17}[\dXx](?!\d)", "[已隐藏]", text)
    return text


def safe_recovery_rows(facts: dict) -> list[list[str]]:
    source_rows = facts.get("safe_table_rows")
    if not isinstance(source_rows, list):
        source_rows = facts.get("table_rows") or []
    result = []
    for row in source_rows:
        if not isinstance(row, dict):
            continue
        cells = [scrub_private_source_value(cell) for cell in (row.get("cells") or [])]
        for index, cell in enumerate(cells[:-1]):
            if is_sensitive_source_label(cell):
                cells[index + 1] = "[已隐藏]"
        if cells:
            result.append(cells)
    return result


def normalized_source_text(value: str) -> str:
    return re.sub(r"\s+", "", strip_indent(str(value or "")))


def source_contains_candidate(facts: dict, value: str) -> bool:
    candidate = normalized_source_text(value)
    if not candidate or candidate == "[已隐藏]":
        return False
    source = normalized_source_text(
        "\n".join(cell for row in safe_recovery_rows(facts) for cell in row)
    )
    return candidate in source


def build_structured_field_recovery_prompt(facts: dict, missing: list[str]) -> str:
    source_rows = json.dumps(safe_recovery_rows(facts), ensure_ascii=False, indent=2)
    return f"""
你负责从出版社选题申报表的脱敏表格文本中找出缺失字段。

只输出严格 JSON，不要解释。不得推测、补写或改写字段值；值必须逐字出现在来源文本中。
仅处理缺失字段：{json.dumps(missing, ensure_ascii=False)}。

输出结构：
{{
  "fields": {{
    "bookName": {{"value": "选题名称原文，没有则为空", "evidence": "对应字段标签"}},
    "authorName": {{"value": "第一作者姓名原文，没有则为空", "evidence": "对应字段标签"}}
  }}
}}

【脱敏表格文本】
{source_rows}
""".strip()


def recover_missing_structured_fields(facts: dict) -> dict:
    recovered = deepcopy(facts if isinstance(facts, dict) else {})
    recovered["canonical_fields"] = dict(recovered.get("canonical_fields") or {})
    missing = missing_core_fields(recovered)
    if not missing:
        return recovered

    runtime_log(f"structured-field fallback start missing={','.join(missing)}")
    try:
        response = call_fast_model(build_structured_field_recovery_prompt(recovered, missing))
    except ModelServiceError as exc:
        runtime_log(
            "structured-field fallback unavailable "
            f"missing={','.join(missing)} type={type(exc).__name__}"
        )
        return recovered

    response_fields = response.get("fields") if isinstance(response, dict) else {}
    if not isinstance(response_fields, dict):
        response_fields = {}
    accepted = []
    for canonical_name in missing:
        model_key = RECOVERY_MODEL_KEYS[canonical_name]
        item = response_fields.get(model_key) or {}
        if not isinstance(item, dict):
            continue
        candidate = strip_indent(str(item.get("value", "")))
        if canonical_name == "姓名":
            candidate = clean_person_name(candidate)
        elif candidate in {"选题名", "选题名称", "教材名称", "书名", "待定"}:
            candidate = ""
        if not candidate or not source_contains_candidate(recovered, candidate):
            continue
        recovered["canonical_fields"][canonical_name] = candidate
        recovered.setdefault("recovered_fields", {})[canonical_name] = "model_grounded"
        accepted.append(canonical_name)

    remaining = missing_core_fields(recovered)
    runtime_log(
        "structured-field fallback complete "
        f"accepted={','.join(accepted) or 'none'} remaining={','.join(remaining) or 'none'}"
    )
    return recovered


def normalize_compare_sources(items) -> list[dict]:
    sources = []
    seen = set()
    for item in items if isinstance(items, list) else []:
        if not isinstance(item, dict):
            continue
        title = strip_indent(str(item.get("title", "") or ""))
        url = strip_indent(str(item.get("url", "") or ""))
        identity = book_title_identity(title)
        if not title or not url.startswith(("http://", "https://")) or not identity or identity in seen:
            continue
        seen.add(identity)
        sources.append(
            {
                "title": title,
                "url": url,
                "author": strip_indent(str(item.get("author", "") or "")),
                "publisher": strip_indent(str(item.get("publisher", "") or "")),
                "year": strip_indent(str(item.get("year", "") or "")),
            }
        )
        if len(sources) >= 2:
            break
    return sources


def normalize_bpm_fields(value: dict | None) -> dict:
    fields = value if isinstance(value, dict) else {}
    return {
        "brief": strip_indent(str(fields.get("brief", "") or "")),
        "reader": strip_indent(str(fields.get("reader", "") or "")),
        "feature": strip_indent(str(fields.get("feature", "") or "")),
        "compare": strip_indent(str(fields.get("compare", "") or "")),
        "compareSources": normalize_compare_sources(
            fields.get("compareSources") or fields.get("compare_sources") or []
        ),
    }


def normalize_generated(result: dict) -> dict:
    sections = []
    section_map = {section.get("key"): section for section in result.get("sections", [])}
    for key, title in SECTION_TITLES.items():
        source = section_map.get(key, {})
        text = FIXED_PROFIT_SECTION_TEXT if key == "profit" else ensure_indent(source.get("text", ""))
        sections.append(
            {
                "key": key,
                "title": source.get("title") or title,
                "text": text,
                "confirmed": False,
            }
        )

    raw_scores = result.get("scores") or []
    score_map = {item[0]: item for item in raw_scores if isinstance(item, list) and len(item) >= 3}
    max_scores = {
        "选题内容": 35,
        "作者情况": 10,
        "策划过程与可行性": 5,
        "获奖潜质": 5,
        "成本与盈利估算": 35,
        "市场定位与营销": 10,
    }
    scores = []
    for name, max_score in max_scores.items():
        value = score_map.get(name, [name, max_score, 0])[2]
        try:
            value = max(0, min(max_score, int(value)))
        except Exception:
            value = 0
        scores.append([name, max_score, value])
    total = sum(item[2] for item in scores)
    bpm_classification = result.get("bpmClassification") or result.get("bpm_classification") or {}
    bpm_fields = normalize_bpm_fields(result.get("bpmFields") or result.get("bpm_fields"))
    author_maintenance = result.get("authorMaintenance") or result.get("author_maintenance") or {}
    return {
        "title": result.get("title") or "选题策划报告",
        "sections": sections,
        "scores": scores,
        "total": total,
        "pending_questions": result.get("pending_questions") or [],
        "bpmClassification": {
            "class3": str(bpm_classification.get("class3", "") or "").strip(),
            "class4": str(bpm_classification.get("class4", "") or "").strip(),
        },
        "bpmFields": bpm_fields,
        "authorMaintenance": {
            "bio": str(author_maintenance.get("bio", "") or "").strip() if isinstance(author_maintenance, dict) else "",
        },
    }


def validate_bpm_field_sources(fields: dict, book_context: list[dict]) -> dict:
    normalized = normalize_bpm_fields(fields)
    allowed = {
        str(item.get("url", "") or ""): item
        for item in book_context
        if isinstance(item, dict) and str(item.get("url", "") or "").startswith(("http://", "https://"))
    }
    verified = []
    for source in normalized["compareSources"]:
        context_item = allowed.get(source["url"])
        if context_item is None:
            continue
        verified.append(source)
    normalized["compareSources"] = verified[:2]
    if len(normalized["compareSources"]) != 2:
        normalized["compare"] = ""
        normalized["compareSources"] = []
    return normalized


def build_dedicated_bpm_fields_prompt(
    facts: dict, report: dict, book_context: list[dict]
) -> str:
    source = compact_facts(facts)
    source["book_comparison_search_context"] = book_context
    report_sections = [
        {
            "key": str(section.get("key", "") or ""),
            "text": strip_indent(str(section.get("text", "") or "")),
        }
        for section in report.get("sections", [])
        if isinstance(section, dict)
    ]
    return f"""
你是电子工业出版社教育出版板块的资深策划编辑。请根据申报表事实、已确认的策划报告和联网检索候选图书，单独生成 BPM 所需的四项内容。只输出严格 JSON。

规则：
1. `brief`：250-400 字，说明本书讲什么、内容组织逻辑和读者能够获得的能力，不得直接复制策划报告第一部分。
2. `reader`：只写一句“本书适合大学XXX专业本科生、研究生，以及从事XXX相关工作的人员使用。”
3. `feature`：250-400 字，说明内容体系、教材组织、实践教学和行业前沿价值，不写编辑工作建议。
4. `compare`：只能从候选图书中恰好选择 2 本，逐本写书名、作者、出版社、年份、优点和相较本选题的覆盖差异，末尾总结本选题差异化定位。不得编造销量、排名和质量评价。
5. `compareSources`：恰好保存上述 2 本书的 `title`、`url`、`author`、`publisher`、`year`，其中 `url` 必须原样取自候选图书。
6. 候选图书不足 2 本时，`compare` 留空且 `compareSources` 输出空数组，不得凭模型记忆补书。
7. 信息缺失时直接略去，不写“申报表没有”“尚未提供”等缺项说明。

输出结构：
{{
  "bpmFields": {{
    "brief": "...",
    "reader": "...",
    "feature": "...",
    "compare": "...",
    "compareSources": [
      {{"title": "...", "url": "...", "author": "...", "publisher": "...", "year": "..."}},
      {{"title": "...", "url": "...", "author": "...", "publisher": "...", "year": "..."}}
    ]
  }}
}}

【申报表事实及候选图书】
{json.dumps(source, ensure_ascii=False, indent=2)}

【已确认策划报告】
{json.dumps(report_sections, ensure_ascii=False, indent=2)}
""".strip()


def generate_dedicated_bpm_fields(facts: dict, report: dict) -> dict:
    book_context = book_comparison_search_context(facts)
    prompt = build_dedicated_bpm_fields_prompt(facts, report, book_context)
    generated = call_complete_report_model(prompt)
    raw_fields = generated.get("bpmFields") or generated.get("bpm_fields") or generated
    return validate_bpm_field_sources(raw_fields, book_context)


def parse_score(value: str, default: int = 0) -> int:
    match = re.search(r"\d+", str(value or ""))
    return int(match.group(0)) if match else default


def docx_unique_row_text(row) -> list[str]:
    seen = []
    cells = []
    for cell in row.cells:
        if cell._tc in seen:
            continue
        seen.append(cell._tc)
        cells.append(strip_indent(cell.text))
    return cells


def extract_report_from_template_table(doc: Document) -> dict:
    if not doc.tables:
        return {}
    table = doc.tables[0]
    if len(table.rows) < 13:
        return {}
    section_keys = ["content", "author", "feasibility", "award", "profit", "marketing"]
    sections = []
    scores = []
    for offset, key in enumerate(section_keys, start=7):
        cells = docx_unique_row_text(table.rows[offset])
        if len(cells) < 3:
            continue
        section_name = SCORE_NAMES[key]
        max_score = {
            "content": 35,
            "author": 10,
            "feasibility": 5,
            "award": 5,
            "profit": 35,
            "marketing": 10,
        }[key]
        text = cells[2] if len(cells) >= 3 else ""
        score = parse_score(cells[3] if len(cells) >= 4 else "", 0)
        sections.append(
            {
                "key": key,
                "title": SECTION_TITLES[key],
                "text": ensure_indent(text),
                "confirmed": True,
            }
        )
        scores.append([section_name, max_score, max(0, min(max_score, score))])
    if len(sections) != 6:
        return {}
    title = ""
    editor_name = ""
    try:
        title = docx_unique_row_text(table.rows[2])[1]
    except Exception:
        title = ""
    try:
        editor_name = docx_unique_row_text(table.rows[3])[1]
    except Exception:
        editor_name = ""
    total = sum(item[2] for item in scores)
    return {
        "title": safe_name(title or "选题策划报告"),
        "editorName": strip_indent(editor_name),
        "sections": sections,
        "scores": scores,
        "total": total,
        "pending_questions": [],
    }


def extract_report_from_plain_text(doc: Document) -> dict:
    chunks = []
    for paragraph in doc.paragraphs:
        text = strip_indent(paragraph.text)
        if text:
            chunks.append(text)
    for table in doc.tables:
        for row in table.rows:
            for cell in docx_unique_row_text(row):
                if cell:
                    chunks.append(cell)
    text = "\n".join(chunks)
    if not text:
        return {}
    title_match = re.search(r"《([^》]+)》?选题策划报告|选题名称[:：\s]*([^\n]+)", text)
    title = title_match.group(1) or title_match.group(2) if title_match else "选题策划报告"
    headings = [
        ("content", r"一[、.．]\s*选题内容"),
        ("author", r"二[、.．]\s*作者情况"),
        ("feasibility", r"三[、.．]\s*策划过程与可行性"),
        ("award", r"四[、.．]\s*获奖潜质"),
        ("profit", r"五[、.．]\s*成本与盈利估算"),
        ("marketing", r"六[、.．]\s*市场定位与营销"),
    ]
    positions = []
    for key, pattern in headings:
        match = re.search(pattern, text)
        if match:
            positions.append((match.start(), match.end(), key))
    positions.sort()
    if len(positions) < 6:
        return {}
    sections = []
    scores = []
    max_scores = {
        "content": 35,
        "author": 10,
        "feasibility": 5,
        "award": 5,
        "profit": 35,
        "marketing": 10,
    }
    defaults = {
        "content": 30,
        "author": 4,
        "feasibility": 5,
        "award": 0,
        "profit": 17,
        "marketing": 9,
    }
    for index, (_, end, key) in enumerate(positions):
        next_start = positions[index + 1][0] if index + 1 < len(positions) else len(text)
        body = text[end:next_start].strip()
        score_match = re.search(r"(?:自评分|建议分|评分)[:：\s]*(\d+)", body)
        score = parse_score(score_match.group(1), defaults[key]) if score_match else defaults[key]
        body = re.sub(r"(?:自评分|建议分|评分)[:：\s]*\d+", "", body).strip()
        sections.append({"key": key, "title": SECTION_TITLES[key], "text": ensure_indent(body), "confirmed": True})
        scores.append([SCORE_NAMES[key], max_scores[key], max(0, min(max_scores[key], score))])
    return {
        "title": safe_name(title or "选题策划报告"),
        "editorName": "",
        "sections": sections,
        "scores": scores,
        "total": sum(item[2] for item in scores),
        "pending_questions": [],
    }


def extract_planning_report_docx(path: Path) -> dict:
    doc = Document(str(path))
    extracted = extract_report_from_template_table(doc) or extract_report_from_plain_text(doc)
    if not extracted:
        raise ValueError("没有从选题策划报告中识别到一到六部分内容，请确认上传的是选题策划报告 DOCX。")
    section_map = {section["key"]: section for section in extracted["sections"]}
    extracted["bpmFields"] = normalize_bpm_fields({})
    extracted["authorMaintenance"] = {"bio": strip_indent(section_map.get("author", {}).get("text", ""))}
    return extracted


def import_bpm_sources(application_bytes: bytes, application_name: str, report_bytes: bytes, report_name: str) -> dict:
    for filename in [application_name, report_name]:
        if Path(filename or "").suffix.lower() != ".docx":
            raise ValueError("请上传 .docx 格式的选题申报表和选题策划报告。")
    application_path = None
    report_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
            tmp.write(application_bytes)
            application_path = Path(tmp.name)
        with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
            tmp.write(report_bytes)
            report_path = Path(tmp.name)
        extractor = load_extractor()
        facts = extractor.build_payload(application_path, include_sensitive=True)
        facts = recover_missing_structured_fields(facts)
        report = extract_planning_report_docx(report_path)
        normalized = normalize_generated(report)
        normalized["title"] = first_non_empty(
            field_value(facts, "教材名称", "选题名称", "选题名", "书名"),
            report.get("title", ""),
            normalized["title"],
        )
        for section in normalized["sections"]:
            section["confirmed"] = True
        normalized["bpmFields"] = generate_dedicated_bpm_fields(facts, normalized)
        normalized["authorMaintenance"] = report.get("authorMaintenance", normalized.get("authorMaintenance", {}))
        normalized["bpmTopic"] = build_bpm_topic(facts, normalized)
        normalized["sourceMode"] = "application-plus-report"
        return normalized
    finally:
        for path_obj in [application_path, report_path]:
            if path_obj:
                try:
                    path_obj.unlink()
                except OSError:
                    pass


def generate_report_from_upload(file_bytes: bytes, filename: str) -> dict:
    suffix = Path(filename or "application.docx").suffix or ".docx"
    if suffix.lower() != ".docx":
        raise ValueError("请上传 .docx 格式的选题申报表，暂不支持 .doc 或其他文件格式。")
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp.write(file_bytes)
        tmp_path = Path(tmp.name)
    try:
        extractor = load_extractor()
        facts = extractor.build_payload(tmp_path, include_sensitive=False)
        facts = recover_missing_structured_fields(facts)
        book_context = book_comparison_search_context(facts)
        prompt = build_prompt(facts, book_context)
        result = call_complete_report_model(prompt)
        result = revise_report_tone_if_needed(facts, result)
        result["bpmFields"] = validate_bpm_field_sources(
            result.get("bpmFields") or result.get("bpm_fields") or {},
            book_context,
        )
        normalized = normalize_generated(result)
        normalized["title"] = first_non_empty(
            field_value(facts, "教材名称", "选题名称", "选题名", "书名"),
            normalized["title"],
        )
        normalized["bpmTopic"] = build_bpm_topic(facts, normalized)
        return normalized
    finally:
        try:
            tmp_path.unlink()
        except OSError:
            pass


def mask_mail_identity(value: str) -> str | None:
    value = str(value or "").strip()
    if not value:
        return None
    if "@" in value:
        local, domain = value.split("@", 1)
        visible = local[:1]
        return f"{visible}{'*' * max(3, len(local) - 1)}@{domain}"
    if len(value) <= 2:
        return "*" * len(value)
    return f"{value[0]}{'*' * max(3, len(value) - 2)}{value[-1]}"


def public_smtp_status(user_id: int) -> dict:
    settings = APP_STORE.get_integration_secret(user_id, "smtp")
    if settings is None:
        return {"configured": False, "accountMasked": None}
    normalized = normalize_smtp_settings(settings)
    return {
        "configured": True,
        "accountMasked": mask_mail_identity(normalized["account"]),
        "fromAddressMasked": mask_mail_identity(normalized["fromAddress"]),
        "host": normalized["host"],
        "port": normalized["port"],
        "security": normalized["security"],
        "fromName": normalized["fromName"],
    }


def safe_mail_error(error: Exception, settings: dict | None = None) -> str:
    known_secrets = []
    if isinstance(settings, dict):
        password = settings.get("password")
        if isinstance(password, str) and password:
            known_secrets.append(password)
    summary = APP_STORE.sanitize_job_data(
        f"{type(error).__name__}: {error}", known_secrets=known_secrets
    )
    return str(summary).strip()[:500] or "邮件发送失败"


def normalize_mail_recipients(recipients) -> list[dict]:
    if not isinstance(recipients, list):
        raise ValueError("收件人名单必须是数组")
    if len(recipients) > MAX_MAIL_RECIPIENTS:
        raise ValueError(f"单次最多发送 {MAX_MAIL_RECIPIENTS} 封邮件")
    normalized = []
    seen = set()
    for recipient in recipients:
        if not isinstance(recipient, dict):
            raise ValueError("收件人信息必须是对象")
        record = {
            str(key): "" if value is None else str(value).strip()
            for key, value in recipient.items()
            if isinstance(key, str)
        }
        record["name"] = str(recipient.get("name") or "").strip()
        record["email"] = str(recipient.get("email") or "").strip()
        email_key = record["email"].casefold()
        if email_key and email_key in seen:
            raise ValueError(f"收件人邮箱重复：{record['email']}")
        if email_key:
            seen.add(email_key)
        normalized.append(record)
    errors = validate_mail_compose("待生成主题", "待生成正文", normalized)
    if errors:
        raise ValueError("；".join(dict.fromkeys(errors)))
    return normalized


def render_mail_deliveries(subject: str, body: str, recipients) -> list[dict]:
    if not isinstance(subject, str) or not isinstance(body, str):
        raise ValueError("邮件主题和正文必须是文本")
    normalized = normalize_mail_recipients(recipients)
    errors = validate_mail_compose(subject, body, normalized)
    if errors:
        raise ValueError("；".join(dict.fromkeys(errors)))
    rendered = []
    for recipient in normalized:
        variables = dict(recipient)
        variables.setdefault("姓名", recipient["name"])
        variables.setdefault("邮箱", recipient["email"])
        rendered_subject = render_mail_text(subject, variables).strip()
        rendered_body = render_mail_text(body, variables)
        rendered_errors = validate_mail_compose(
            rendered_subject, rendered_body, [recipient]
        )
        if rendered_errors:
            raise ValueError("；".join(dict.fromkeys(rendered_errors)))
        rendered.append(
            {
                "name": recipient["name"],
                "email": recipient["email"],
                "subject": rendered_subject,
                "body": rendered_body,
            }
        )
    return rendered


def mail_request_fingerprint(operation: str, payload: dict) -> str:
    canonical = json.dumps(
        {"operation": operation, **payload},
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(canonical).hexdigest()


def fail_unfinished_mail_batch(
    user_id: int, batch_id: str, error: Exception, settings: dict | None = None
) -> None:
    batch = APP_STORE.get_mail_batch(user_id, batch_id)
    if batch is None or batch["status"] != "running":
        return
    error_summary = safe_mail_error(error, settings)
    for delivery in batch["deliveries"]:
        if delivery["status"] == "queued":
            APP_STORE.update_mail_delivery(
                user_id, batch_id, delivery["id"], "sending"
            )
            APP_STORE.update_mail_delivery(
                user_id, batch_id, delivery["id"], "failed", error_summary
            )
        elif delivery["status"] == "sending":
            APP_STORE.update_mail_delivery(
                user_id, batch_id, delivery["id"], "failed", error_summary
            )
    APP_STORE.finish_mail_batch(user_id, batch_id)


def process_mail_batch(batch_id: str, user_id: int):
    settings = {}
    batch = None
    try:
        settings = APP_STORE.get_integration_secret(user_id, "smtp") or {}
        batch = APP_STORE.start_mail_batch(user_id, batch_id)
        if batch is None:
            return
        normalized_settings = normalize_smtp_settings(settings)
        settings.clear()
        settings.update(normalized_settings)
        for delivery in batch["deliveries"]:
            claimed = APP_STORE.update_mail_delivery(
                user_id, batch_id, delivery["id"], "sending"
            )
            if not claimed:
                continue
            try:
                send_smtp_message(settings, delivery)
                APP_STORE.update_mail_delivery(
                    user_id, batch_id, delivery["id"], "succeeded"
                )
            except Exception as error:
                APP_STORE.update_mail_delivery(
                    user_id,
                    batch_id,
                    delivery["id"],
                    "failed",
                    safe_mail_error(error, settings),
                )
        APP_STORE.finish_mail_batch(user_id, batch_id)
    except Exception as error:
        if batch is not None:
            fail_unfinished_mail_batch(user_id, batch_id, error, settings)
        runtime_log(
            f"mail-batch failure batch={batch_id} type={type(error).__name__}"
        )
    finally:
        if isinstance(settings, dict):
            settings.clear()


def schedule_mail_batch(batch_id: str, user_id: int) -> None:
    threading.Thread(
        target=process_mail_batch, args=(batch_id, user_id), daemon=True
    ).start()


class Handler(BaseHTTPRequestHandler):
    def handle_one_request(self):
        try:
            super().handle_one_request()
        except RequestHandled:
            return

    def end_headers(self):
        self.send_header("Cache-Control", "no-store, no-cache, must-revalidate, max-age=0")
        self.send_header("Pragma", "no-cache")
        self.send_header("Expires", "0")
        super().end_headers()

    def send_json(self, status: int, payload: dict, session_cookie: str | None = None):
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        if session_cookie is not None:
            self.send_header("Set-Cookie", session_cookie)
        self.end_headers()
        self.wfile.write(body)

    def serve_static(self, path: str, include_body: bool = True):
        asset = STATIC_ASSETS.get(path)
        if asset is None:
            self.send_error(404)
            return
        relative_path, content_type = asset
        try:
            data = (ROOT / relative_path).read_bytes()
        except OSError:
            self.send_error(404)
            return
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        if include_body:
            self.wfile.write(data)

    def read_json(self) -> dict:
        try:
            length = int(self.headers.get("Content-Length", "0"))
        except (TypeError, ValueError) as error:
            raise ValueError("Content-Length 无效。") from error
        if length <= 0:
            raise ValueError("JSON 请求体不能为空。")
        if length > MAX_JSON_BODY_BYTES:
            self.close_connection = True
            self.send_json(
                413,
                {
                    "error": "JSON 请求体不得超过 2 MB。",
                    "code": "REQUEST_BODY_TOO_LARGE",
                },
            )
            raise RequestHandled()

        previous_timeout = self.connection.gettimeout()
        self.connection.settimeout(JSON_READ_TIMEOUT_SECONDS)
        try:
            raw = self.rfile.read(length)
        except (TimeoutError, socket.timeout):
            self.close_connection = True
            self.send_json(
                408,
                {"error": "读取请求体超时。", "code": "REQUEST_BODY_TIMEOUT"},
            )
            raise RequestHandled()
        finally:
            self.connection.settimeout(previous_timeout)
        if len(raw) != length:
            raise ValueError("JSON 请求体不完整。")
        payload = json.loads(raw.decode("utf-8"))
        if not isinstance(payload, dict):
            raise ValueError("JSON 请求体必须是对象。")
        return payload

    @staticmethod
    def public_user(user: dict) -> dict:
        return {
            "id": user["id"],
            "username": user["username"],
            "displayName": user["display_name"],
        }

    @staticmethod
    def session_cookie(token: str, max_age: int) -> str:
        cookie = SimpleCookie()
        cookie[SESSION_COOKIE_NAME] = token
        morsel = cookie[SESSION_COOKIE_NAME]
        morsel["path"] = "/"
        morsel["httponly"] = True
        morsel["samesite"] = "Lax"
        morsel["max-age"] = max_age
        return morsel.OutputString()

    def current_user(self) -> dict | None:
        try:
            cookies = SimpleCookie(self.headers.get("Cookie", ""))
        except CookieError:
            return None
        morsel = cookies.get(SESSION_COOKIE_NAME)
        return None if morsel is None else APP_STORE.get_user_for_session(morsel.value)

    def require_user(self) -> dict | None:
        user = self.current_user()
        if user:
            return user
        self.send_json(401, {"error": "请先登录。", "code": "AUTH_REQUIRED"})
        return None

    def handle_auth_register(self):
        try:
            payload = self.read_json()
            user = APP_STORE.register_user(
                payload.get("username", ""),
                payload.get("password", ""),
                payload.get("displayName", ""),
            )
            token = APP_STORE.create_session(user["id"])
            self.send_json(201, {"user": self.public_user(user)}, self.session_cookie(token, 604800))
        except (TypeError, ValueError, json.JSONDecodeError) as error:
            self.send_json(400, {"error": str(error), "code": "AUTH_INVALID_INPUT"})

    def handle_auth_login(self):
        try:
            payload = self.read_json()
        except (TypeError, ValueError, json.JSONDecodeError) as error:
            self.send_json(400, {"error": str(error), "code": "AUTH_INVALID_INPUT"})
            return
        user = APP_STORE.authenticate_user(payload.get("username", ""), payload.get("password", ""))
        if user is None:
            self.send_json(401, {"error": "用户名或密码错误。", "code": "AUTH_INVALID"})
            return
        token = APP_STORE.create_session(user["id"])
        self.send_json(200, {"user": self.public_user(user)}, self.session_cookie(token, 604800))

    def handle_auth_logout(self):
        if not self.require_user():
            return
        try:
            cookies = SimpleCookie(self.headers.get("Cookie", ""))
            morsel = cookies.get(SESSION_COOKIE_NAME)
            if morsel is not None:
                APP_STORE.delete_session(morsel.value)
        except CookieError:
            pass
        self.send_json(200, {"ok": True}, self.session_cookie("", 0))

    def handle_bpm_integration_get(self, user: dict):
        try:
            self.send_json(200, APP_STORE.get_integration_status(user["id"], "phei_bpm"))
        except ValueError as error:
            self.send_json(400, {"error": str(error), "code": "INTEGRATION_CONFIGURATION_ERROR"})

    def handle_bpm_integration_put(self, user: dict):
        try:
            payload = self.read_json()
            account = payload.get("account")
            password = payload.get("password")
            if not isinstance(account, str) or not account.strip():
                raise ValueError("BPM 账号不能为空。")
            if not isinstance(password, str) or not password.strip():
                raise ValueError("BPM 密码不能为空。")
            APP_STORE.put_integration_credentials(
                user["id"], "phei_bpm", account.strip(), password
            )
            self.send_json(200, APP_STORE.get_integration_status(user["id"], "phei_bpm"))
        except (TypeError, ValueError, json.JSONDecodeError) as error:
            self.send_json(400, {"error": str(error), "code": "INTEGRATION_INVALID_INPUT"})

    def handle_bpm_integration_delete(self, user: dict):
        try:
            APP_STORE.delete_integration_credentials(user["id"], "phei_bpm")
            self.send_json(200, {"configured": False, "accountMasked": None})
        except ValueError as error:
            self.send_json(400, {"error": str(error), "code": "INTEGRATION_CONFIGURATION_ERROR"})

    def handle_smtp_integration_get(self, user: dict):
        try:
            self.send_json(200, public_smtp_status(user["id"]))
        except ValueError as error:
            self.send_json(
                400, {"error": str(error), "code": "SMTP_CONFIGURATION_ERROR"}
            )

    def handle_smtp_integration_put(self, user: dict):
        try:
            settings = normalize_smtp_settings(self.read_json())
            APP_STORE.put_integration_secret(user["id"], "smtp", settings)
            self.send_json(200, public_smtp_status(user["id"]))
        except (TypeError, ValueError, json.JSONDecodeError) as error:
            self.send_json(400, {"error": str(error), "code": "SMTP_INVALID_INPUT"})

    def handle_smtp_integration_delete(self, user: dict):
        try:
            APP_STORE.delete_integration_credentials(user["id"], "smtp")
            self.send_json(200, {"configured": False, "accountMasked": None})
        except ValueError as error:
            self.send_json(
                400, {"error": str(error), "code": "SMTP_CONFIGURATION_ERROR"}
            )

    def handle_smtp_test(self, user: dict):
        settings = {}
        try:
            settings = APP_STORE.get_integration_secret(user["id"], "smtp") or {}
            if not settings:
                self.send_json(
                    400,
                    {"error": "请先保存 SMTP 配置。", "code": "SMTP_NOT_CONFIGURED"},
                )
                return
            test_smtp_connection(settings)
            self.send_json(200, {"ok": True})
        except Exception as error:
            self.send_json(
                400,
                {"error": safe_mail_error(error, settings), "code": "SMTP_TEST_FAILED"},
            )
        finally:
            if isinstance(settings, dict):
                settings.clear()

    def handle_mail_templates_list(self, user: dict):
        templates = APP_STORE.ensure_default_mail_templates(
            user["id"], DEFAULT_MAIL_TEMPLATES
        )
        self.send_json(200, {"templates": templates})

    def handle_mail_template_create(self, user: dict):
        try:
            payload = self.read_json()
            template = APP_STORE.create_mail_template(
                user["id"],
                payload.get("name"),
                payload.get("subject"),
                payload.get("body"),
            )
            self.send_json(201, {"template": template})
        except (TypeError, ValueError, json.JSONDecodeError) as error:
            self.send_json(
                400, {"error": str(error), "code": "MAIL_TEMPLATE_INVALID"}
            )

    def handle_mail_template_update(self, user: dict, template_id: str):
        try:
            payload = self.read_json()
            template = APP_STORE.update_mail_template(
                user["id"],
                template_id,
                payload.get("name"),
                payload.get("subject"),
                payload.get("body"),
            )
            if template is None:
                self.send_json(
                    404,
                    {"error": "未找到邮件模板。", "code": "MAIL_TEMPLATE_NOT_FOUND"},
                )
                return
            self.send_json(200, {"template": template})
        except (TypeError, ValueError, json.JSONDecodeError) as error:
            self.send_json(
                400, {"error": str(error), "code": "MAIL_TEMPLATE_INVALID"}
            )

    def handle_mail_template_delete(self, user: dict, template_id: str):
        if not APP_STORE.delete_mail_template(user["id"], template_id):
            self.send_json(
                404,
                {"error": "未找到邮件模板。", "code": "MAIL_TEMPLATE_NOT_FOUND"},
            )
            return
        self.send_json(200, {"ok": True})

    def handle_mail_draft_get(self, user: dict):
        self.send_json(200, {"draft": APP_STORE.get_mail_draft(user["id"])})

    def handle_mail_draft_put(self, user: dict):
        try:
            payload = self.read_json()
            APP_STORE.put_mail_draft(user["id"], payload)
            self.send_json(200, {"draft": APP_STORE.get_mail_draft(user["id"])})
        except (TypeError, ValueError, json.JSONDecodeError) as error:
            self.send_json(400, {"error": str(error), "code": "MAIL_DRAFT_INVALID"})

    def handle_mail_recipient_parse(self, user: dict):
        try:
            content_length = int(self.headers.get("Content-Length", "0"))
            if content_length <= 0 or content_length > MAX_MAIL_RECIPIENT_FILE_BYTES:
                raise ValueError("名单文件不能为空且不得超过 10 MB")
            form = read_multipart_form(
                fp=self.rfile,
                headers=self.headers,
                environ={
                    "REQUEST_METHOD": "POST",
                    "CONTENT_TYPE": self.headers.get("Content-Type", ""),
                    "CONTENT_LENGTH": str(content_length),
                },
            )
            file_item = form["file"] if "file" in form else None
            if file_item is None or not getattr(file_item, "file", None):
                raise ValueError("请上传 CSV 或 XLSX 名单文件")
            filename = getattr(file_item, "filename", "recipients.csv")
            result = parse_recipient_file(file_item.file.read(), filename)
            self.send_json(200, result)
        except Exception as error:
            self.send_json(
                400,
                {"error": str(error), "code": "MAIL_RECIPIENT_FILE_INVALID"},
            )

    def handle_mail_preview(self, user: dict):
        try:
            payload = self.read_json()
            preview = render_mail_deliveries(
                payload.get("subject"), payload.get("body"), payload.get("recipients")
            )
            self.send_json(200, {"preview": preview})
        except (TypeError, ValueError, json.JSONDecodeError) as error:
            self.send_json(
                400, {"error": str(error), "code": "MAIL_COMPOSE_INVALID"}
            )

    def handle_mail_batch_list(self, user: dict, parsed_url):
        try:
            query = parse_qs(parsed_url.query, keep_blank_values=False)
            limit = int(query.get("limit", ["30"])[0])
            if not 1 <= limit <= 100:
                raise ValueError("查询数量必须在 1—100 之间")
            self.send_json(
                200, {"batches": APP_STORE.list_mail_batches(user["id"], limit=limit)}
            )
        except (TypeError, ValueError) as error:
            self.send_json(400, {"error": str(error), "code": "MAIL_BATCH_INVALID"})

    def handle_mail_batch_get(self, user: dict, batch_id: str):
        batch = APP_STORE.get_mail_batch(user["id"], batch_id)
        if batch is None:
            self.send_json(
                404, {"error": "未找到邮件批次。", "code": "MAIL_BATCH_NOT_FOUND"}
            )
            return
        self.send_json(200, {"batch": batch})

    def require_saved_smtp(self, user: dict) -> dict | None:
        settings = APP_STORE.get_integration_secret(user["id"], "smtp")
        if settings is None:
            self.send_json(
                400,
                {"error": "请先保存 SMTP 配置。", "code": "SMTP_NOT_CONFIGURED"},
            )
            return None
        try:
            return normalize_smtp_settings(settings)
        except ValueError as error:
            self.send_json(
                400, {"error": str(error), "code": "SMTP_CONFIGURATION_ERROR"}
            )
            return None
        finally:
            settings.clear()

    @staticmethod
    def payload_contains_mail_secret(payload: dict) -> bool:
        forbidden = {"password", "smtp", "smtpsettings", "smtpconfig", "authorization"}
        return any(str(key).casefold() in forbidden for key in payload)

    def require_idempotency_key(self) -> str | None:
        request_key = self.headers.get("Idempotency-Key", "").strip()
        if IDEMPOTENCY_KEY_RE.fullmatch(request_key):
            return request_key
        self.send_json(
            400,
            {
                "error": "Idempotency-Key 必须为 8—128 位字母、数字或 . _ : -。",
                "code": "MAIL_IDEMPOTENCY_KEY_INVALID",
            },
        )
        return None

    def handle_mail_batch_create(self, user: dict):
        try:
            request_key = self.require_idempotency_key()
            if request_key is None:
                return
            payload = self.read_json()
            if self.payload_contains_mail_secret(payload):
                self.send_json(
                    400,
                    {
                        "error": "发送任务不允许从客户端传入 SMTP 密码。",
                        "code": "MAIL_CLIENT_SECRET_FORBIDDEN",
                    },
                )
                return
            if payload.get("confirmed") is not True:
                self.send_json(
                    400,
                    {
                        "error": "请确认收件人数量和邮件内容后再发送。",
                        "code": "MAIL_CONFIRMATION_REQUIRED",
                    },
                )
                return
            settings = self.require_saved_smtp(user)
            if settings is None:
                return
            settings.clear()
            subject = payload.get("subject")
            body = payload.get("body")
            deliveries = render_mail_deliveries(
                subject, body, payload.get("recipients")
            )
            normalized_subject = subject.strip()
            request_fingerprint = mail_request_fingerprint(
                "create",
                {
                    "subject": normalized_subject,
                    "body": body,
                    "deliveries": deliveries,
                },
            )
            batch, created = APP_STORE.create_mail_batch_idempotent(
                user["id"],
                normalized_subject,
                body,
                deliveries,
                request_key,
                request_fingerprint,
            )
            template_id = payload.get("templateId")
            if created and isinstance(template_id, str) and template_id:
                APP_STORE.touch_mail_template(user["id"], template_id)
            if created:
                schedule_mail_batch(batch["id"], user["id"])
            self.send_json(201 if created else 200, {"batch": batch, "created": created})
        except IdempotencyConflict as error:
            self.send_json(
                409, {"error": str(error), "code": "MAIL_IDEMPOTENCY_CONFLICT"}
            )
        except (TypeError, ValueError, json.JSONDecodeError) as error:
            self.send_json(400, {"error": str(error), "code": "MAIL_BATCH_INVALID"})

    def handle_mail_batch_retry(self, user: dict, batch_id: str):
        try:
            request_key = self.require_idempotency_key()
            if request_key is None:
                return
            payload = self.read_json()
            if payload.get("confirmed") is not True:
                self.send_json(
                    400,
                    {
                        "error": "请确认只重试失败邮件。",
                        "code": "MAIL_CONFIRMATION_REQUIRED",
                    },
                )
                return
            settings = self.require_saved_smtp(user)
            if settings is None:
                return
            settings.clear()
            request_fingerprint = mail_request_fingerprint(
                "retry", {"batchId": batch_id}
            )
            batch, created = APP_STORE.create_retry_mail_batch_idempotent(
                user["id"], batch_id, request_key, request_fingerprint
            )
            if created:
                schedule_mail_batch(batch["id"], user["id"])
            self.send_json(201 if created else 200, {"batch": batch, "created": created})
        except IdempotencyConflict as error:
            self.send_json(
                409, {"error": str(error), "code": "MAIL_IDEMPOTENCY_CONFLICT"}
            )
        except json.JSONDecodeError as error:
            self.send_json(400, {"error": str(error), "code": "MAIL_BATCH_INVALID"})
        except ValueError as error:
            if APP_STORE.get_mail_batch(user["id"], batch_id) is None:
                self.send_json(
                    404,
                    {"error": "未找到邮件批次。", "code": "MAIL_BATCH_NOT_FOUND"},
                )
                return
            self.send_json(400, {"error": str(error), "code": "MAIL_BATCH_INVALID"})

    @staticmethod
    def project_api_match(path: str):
        return PROJECT_API_RE.fullmatch(path)

    @staticmethod
    def public_project_file(file_record: dict) -> dict:
        return {
            "id": file_record["id"],
            "projectId": file_record["projectId"],
            "kind": file_record["kind"],
            "originalName": file_record["originalName"],
            "storagePath": file_record["storagePath"],
            "sha256": file_record["sha256"],
            "sizeBytes": file_record["sizeBytes"],
            "createdAt": file_record["createdAt"],
        }

    @staticmethod
    def project_source_file(file_record: dict) -> dict:
        return {
            "fileId": file_record["id"],
            "kind": file_record["kind"],
            "originalName": file_record["originalName"],
            "sha256": file_record["sha256"],
            "sizeBytes": file_record["sizeBytes"],
        }

    def save_project_docx(
        self,
        user: dict,
        project_id: str,
        kind: str,
        filename: str,
        content: bytes,
    ) -> dict:
        file_id = uuid.uuid4().hex
        stored = PROJECT_FILE_STORE.save_docx(
            user["id"], project_id, file_id, content
        )
        try:
            return APP_STORE.add_project_file(
                user["id"],
                project_id,
                file_id,
                kind,
                filename,
                stored.relative_path,
                stored.sha256,
                stored.size_bytes,
            )
        except Exception:
            PROJECT_FILE_STORE.remove(stored.relative_path)
            raise

    def rollback_project_file(self, user: dict, project_id: str, record: dict):
        try:
            APP_STORE.delete_project_file(user["id"], project_id, record["id"])
        finally:
            PROJECT_FILE_STORE.remove(record["storagePath"])

    def persist_project_report(
        self,
        user: dict,
        project: dict,
        result: dict,
        *,
        source_kind: str,
        source_files: dict,
        expected_version: int | None = None,
        revision_reason: str,
    ) -> dict:
        state = project_state_from_report(result, source_kind=source_kind)
        state["sourceFiles"].update(source_files)
        summary = summarize_project_state(state, user["display_name"])
        return APP_STORE.update_project(
            user["id"],
            project["id"],
            project["version"] if expected_version is None else expected_version,
            {
                "title": summary["title"],
                "author_name": summary["authorName"],
                "editor_name": summary["editorName"],
                "report_status": summary["reportStatus"],
                "author_status": summary["authorStatus"],
                "bpm_status": summary["bpmStatus"],
                "state": summary["state"],
            },
            revision_reason=revision_reason,
        )

    def project_payload(self, user: dict, project: dict, *, include_preflight: bool = False):
        payload = {
            "project": project,
            "files": [
                self.public_project_file(item)
                for item in APP_STORE.list_project_files(user["id"], project["id"])
            ],
        }
        if include_preflight:
            try:
                configured = APP_STORE.has_integration_credentials(user["id"], "phei_bpm")
            except ValueError:
                configured = False
            payload["preflight"] = build_bpm_preflight(
                project["state"], bpm_configured=configured
            )
        return payload

    def handle_project_list(self, user: dict, parsed_url):
        query = parse_qs(parsed_url.query, keep_blank_values=False)
        include_archived = query.get("includeArchived", [""])[0].lower() == "true"
        projects = APP_STORE.list_projects(
            user["id"],
            include_archived=include_archived,
            query=query.get("q", [""])[0],
            status=query.get("status", [""])[0],
        )
        self.send_json(200, {"projects": projects})

    def handle_project_create(self, user: dict):
        try:
            payload = self.read_json()
            title = payload.get("title", "")
            if not isinstance(title, str):
                raise ValueError("选题名称必须是文本。")
            state = normalize_project_state(payload.get("state"))
            if title.strip():
                state["bpmTopic"]["bookName"] = title.strip()
                state["fieldSources"].setdefault("bpmTopic.bookName", "user")
            summary = summarize_project_state(state, user["display_name"])
            project = APP_STORE.create_project(
                user["id"],
                summary["title"],
                summary["state"],
                author_name=summary["authorName"],
                editor_name=summary["editorName"],
                report_status=summary["reportStatus"],
                author_status=summary["authorStatus"],
                bpm_status=summary["bpmStatus"],
            )
            self.send_json(201, self.project_payload(user, project))
        except (TypeError, ValueError, json.JSONDecodeError) as error:
            self.send_json(400, {"error": str(error), "code": "PROJECT_INVALID_INPUT"})

    def handle_project_get(self, user: dict, project_id: str):
        project = APP_STORE.get_project(user["id"], project_id)
        if project is None:
            self.send_json(404, {"error": "未找到选题项目。", "code": "PROJECT_NOT_FOUND"})
            return
        self.send_json(200, self.project_payload(user, project))

    def handle_project_update(self, user: dict, project_id: str):
        project = APP_STORE.get_project(user["id"], project_id)
        if project is None:
            self.send_json(404, {"error": "未找到选题项目。", "code": "PROJECT_NOT_FOUND"})
            return
        try:
            payload = self.read_json()
            version = payload.get("version")
            if not isinstance(version, int):
                raise ValueError("项目版本无效。")
            state = normalize_project_state(payload.get("state", project["state"]))
            if "title" in payload:
                if not isinstance(payload["title"], str):
                    raise ValueError("选题名称必须是文本。")
                state["bpmTopic"]["bookName"] = payload["title"].strip()
                state["fieldSources"]["bpmTopic.bookName"] = "user"
            summary = summarize_project_state(state, user["display_name"])
            author_status = summary["authorStatus"]
            if project["authorStatus"] in {"queued", "running", "succeeded", "failed"}:
                author_status = project["authorStatus"]
            bpm_status = summary["bpmStatus"]
            if project["bpmStatus"] in {"queued", "running", "succeeded", "failed"}:
                bpm_status = project["bpmStatus"]
            updated = APP_STORE.update_project(
                user["id"],
                project_id,
                version,
                {
                    "title": summary["title"],
                    "author_name": summary["authorName"],
                    "editor_name": summary["editorName"],
                    "report_status": summary["reportStatus"],
                    "author_status": author_status,
                    "bpm_status": bpm_status,
                    "state": summary["state"],
                },
                revision_reason=payload.get("revisionReason"),
            )
            self.send_json(200, self.project_payload(user, updated))
        except ProjectVersionConflict as error:
            self.send_json(
                409,
                {"error": str(error), "code": "PROJECT_VERSION_CONFLICT"},
            )
        except (TypeError, ValueError, json.JSONDecodeError) as error:
            self.send_json(400, {"error": str(error), "code": "PROJECT_INVALID_INPUT"})

    def handle_project_archive(self, user: dict, project_id: str):
        project = APP_STORE.get_project(user["id"], project_id)
        if project is None:
            self.send_json(404, {"error": "未找到选题项目。", "code": "PROJECT_NOT_FOUND"})
            return
        try:
            payload = self.read_json()
            archived = APP_STORE.archive_project(
                user["id"], project_id, payload.get("version")
            )
            self.send_json(200, self.project_payload(user, archived))
        except ProjectVersionConflict as error:
            self.send_json(
                409,
                {"error": str(error), "code": "PROJECT_VERSION_CONFLICT"},
            )
        except (TypeError, ValueError, json.JSONDecodeError) as error:
            self.send_json(400, {"error": str(error), "code": "PROJECT_INVALID_INPUT"})

    def handle_project_preflight(self, user: dict, project_id: str):
        project = APP_STORE.get_project(user["id"], project_id)
        if project is None:
            self.send_json(404, {"error": "未找到选题项目。", "code": "PROJECT_NOT_FOUND"})
            return
        self.send_json(200, self.project_payload(user, project, include_preflight=True))

    def handle_project_file_upload(self, user: dict, project_id: str):
        project = APP_STORE.get_project(user["id"], project_id)
        if project is None:
            self.send_json(404, {"error": "未找到选题项目。", "code": "PROJECT_NOT_FOUND"})
            return
        stored = None
        try:
            form = read_multipart_form(
                fp=self.rfile,
                headers=self.headers,
                environ={
                    "REQUEST_METHOD": "POST",
                    "CONTENT_TYPE": self.headers.get("Content-Type", ""),
                    "CONTENT_LENGTH": self.headers.get("Content-Length", "0"),
                },
            )
            kind = form.getfirst("kind", "application")
            if isinstance(kind, bytes):
                kind = kind.decode("utf-8")
            file_item = form["file"] if "file" in form else None
            if file_item is None or not getattr(file_item, "file", None):
                raise InvalidProjectFile("请上传 DOCX 文件。")
            filename = getattr(file_item, "filename", "document.docx")
            if Path(filename).suffix.lower() != ".docx":
                raise InvalidProjectFile("仅支持 .docx 文件。")
            content = file_item.file.read()
            file_id = uuid.uuid4().hex
            stored = PROJECT_FILE_STORE.save_docx(
                user["id"], project_id, file_id, content
            )
            record = APP_STORE.add_project_file(
                user["id"],
                project_id,
                file_id,
                kind,
                filename,
                stored.relative_path,
                stored.sha256,
                stored.size_bytes,
            )
            self.send_json(201, {"file": self.public_project_file(record)})
        except InvalidProjectFile as error:
            if stored is not None:
                PROJECT_FILE_STORE.remove(stored.relative_path)
            self.send_json(400, {"error": str(error), "code": "PROJECT_FILE_INVALID"})
        except ValueError as error:
            if stored is not None:
                PROJECT_FILE_STORE.remove(stored.relative_path)
            self.send_json(400, {"error": str(error), "code": "PROJECT_FILE_INVALID"})

    def handle_project_import(self, user: dict):
        project = None
        records = []
        try:
            form = read_multipart_form(
                fp=self.rfile,
                headers=self.headers,
                environ={
                    "REQUEST_METHOD": "POST",
                    "CONTENT_TYPE": self.headers.get("Content-Type", ""),
                    "CONTENT_LENGTH": self.headers.get("Content-Length", "0"),
                },
            )
            application_item = form["application"] if "application" in form else None
            report_item = form["report"] if "report" in form else None
            if application_item is None or not getattr(application_item, "file", None):
                raise ValueError("请上传选题申报表 DOCX。")
            if report_item is None or not getattr(report_item, "file", None):
                raise ValueError("请上传选题策划报告 DOCX。")
            application_bytes = application_item.file.read()
            report_bytes = report_item.file.read()
            application_name = getattr(application_item, "filename", "application.docx")
            report_name = getattr(report_item, "filename", "report.docx")
            result = import_bpm_sources(
                application_bytes,
                application_name,
                report_bytes,
                report_name,
            )
            initial_state = project_state_from_report(result, source_kind="imported")
            initial_summary = summarize_project_state(initial_state, user["display_name"])
            project = APP_STORE.create_project(
                user["id"],
                initial_summary["title"],
                initial_summary["state"],
                author_name=initial_summary["authorName"],
                editor_name=initial_summary["editorName"],
                report_status=initial_summary["reportStatus"],
                author_status=initial_summary["authorStatus"],
                bpm_status=initial_summary["bpmStatus"],
            )
            application_record = self.save_project_docx(
                user,
                project["id"],
                "application",
                application_name,
                application_bytes,
            )
            records.append(application_record)
            report_record = self.save_project_docx(
                user,
                project["id"],
                "confirmed_report",
                report_name,
                report_bytes,
            )
            records.append(report_record)
            project = self.persist_project_report(
                user,
                project,
                result,
                source_kind="imported",
                source_files={
                    "application": self.project_source_file(application_record),
                    "report": self.project_source_file(report_record),
                },
                revision_reason="quick_import",
            )
            response = dict(result)
            response.update(self.project_payload(user, project, include_preflight=True))
            self.send_json(201, response)
        except (InvalidProjectFile, TypeError, ValueError, json.JSONDecodeError) as error:
            if project is not None:
                for record in reversed(records):
                    self.rollback_project_file(user, project["id"], record)
            self.send_json(400, {"error": str(error), "code": "PROJECT_IMPORT_INVALID"})
        except Exception as error:
            if project is not None:
                for record in reversed(records):
                    self.rollback_project_file(user, project["id"], record)
            runtime_log(f"project-import error {type(error).__name__}: {error}")
            self.send_json(500, {"error": str(error), "code": "PROJECT_IMPORT_FAILED"})

    def handle_project_file_download(
        self, user: dict, project_id: str, file_id: str
    ):
        record = APP_STORE.get_project_file(user["id"], project_id, file_id)
        if record is None:
            self.send_json(404, {"error": "未找到项目文件。", "code": "PROJECT_FILE_NOT_FOUND"})
            return
        try:
            data = PROJECT_FILE_STORE.resolve(record["storagePath"]).read_bytes()
        except (InvalidProjectFile, OSError):
            self.send_json(404, {"error": "项目文件已丢失。", "code": "PROJECT_FILE_MISSING"})
            return
        self.send_response(200)
        self.send_header(
            "Content-Type",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.send_header("Content-Length", str(len(data)))
        self.send_header(
            "Content-Disposition",
            f"attachment; filename*=UTF-8''{quote(record['originalName'])}",
        )
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        parsed_url = urlparse(self.path)
        path = unquote(parsed_url.path)
        if path == "/api/health":
            self.send_json(200, {"ok": True, "time": now_iso()})
            return
        if path == "/api/auth/me":
            user = self.require_user()
            if user:
                self.send_json(200, {"user": self.public_user(user)})
            return
        user = None
        if path.startswith("/api/"):
            user = self.require_user()
            if not user:
                return
        if path == "/api/integrations/phei-bpm":
            self.handle_bpm_integration_get(user)
            return
        if path == "/api/integrations/smtp":
            self.handle_smtp_integration_get(user)
            return
        if path == "/api/mail/templates":
            self.handle_mail_templates_list(user)
            return
        if path == "/api/mail/draft":
            self.handle_mail_draft_get(user)
            return
        if path == "/api/mail/batches":
            self.handle_mail_batch_list(user, parsed_url)
            return
        mail_batch_match = MAIL_BATCH_API_RE.fullmatch(path)
        if mail_batch_match and mail_batch_match.group(2) is None:
            self.handle_mail_batch_get(user, mail_batch_match.group(1))
            return
        if path == "/api/bpm-jobs":
            self.send_json(200, {"jobs": APP_STORE.list_jobs(user["id"])})
            return
        if path == "/api/projects":
            self.handle_project_list(user, parsed_url)
            return
        project_match = self.project_api_match(path)
        if project_match:
            project_id, action, file_id = project_match.groups()
            if action is None:
                self.handle_project_get(user, project_id)
                return
            if action == "preflight" and file_id is None:
                self.handle_project_preflight(user, project_id)
                return
            if action == "files" and file_id:
                self.handle_project_file_download(user, project_id, file_id)
                return
        self.serve_static(path)

    def do_HEAD(self):
        path = unquote(self.path.split("?", 1)[0])
        self.serve_static(path, include_body=False)

    def do_PUT(self):
        path = unquote(urlparse(self.path).path)
        user = self.require_user() if path.startswith("/api/") else None
        if path.startswith("/api/") and not user:
            return
        if path == "/api/integrations/phei-bpm":
            self.handle_bpm_integration_put(user)
            return
        if path == "/api/integrations/smtp":
            self.handle_smtp_integration_put(user)
            return
        if path == "/api/mail/draft":
            self.handle_mail_draft_put(user)
            return
        mail_template_match = MAIL_TEMPLATE_API_RE.fullmatch(path)
        if mail_template_match:
            self.handle_mail_template_update(user, mail_template_match.group(1))
            return
        project_match = self.project_api_match(path)
        if project_match and project_match.group(2) is None:
            self.handle_project_update(user, project_match.group(1))
            return
        self.send_error(404)

    def do_DELETE(self):
        path = unquote(urlparse(self.path).path)
        user = self.require_user() if path.startswith("/api/") else None
        if path.startswith("/api/") and not user:
            return
        if path == "/api/integrations/phei-bpm":
            self.handle_bpm_integration_delete(user)
            return
        if path == "/api/integrations/smtp":
            self.handle_smtp_integration_delete(user)
            return
        mail_template_match = MAIL_TEMPLATE_API_RE.fullmatch(path)
        if mail_template_match:
            self.handle_mail_template_delete(user, mail_template_match.group(1))
            return
        self.send_error(404)

    def do_POST(self):
        path = unquote(urlparse(self.path).path)
        if path == "/api/auth/register":
            self.handle_auth_register()
            return
        if path == "/api/auth/login":
            self.handle_auth_login()
            return
        if path == "/api/auth/logout":
            self.handle_auth_logout()
            return
        user = None
        if path.startswith("/api/"):
            user = self.require_user()
            if not user:
                return
        if path == "/api/integrations/smtp/test":
            self.handle_smtp_test(user)
            return
        if path == "/api/mail/templates":
            self.handle_mail_template_create(user)
            return
        if path == "/api/mail/recipients/parse":
            self.handle_mail_recipient_parse(user)
            return
        if path == "/api/mail/preview":
            self.handle_mail_preview(user)
            return
        if path == "/api/mail/batches":
            self.handle_mail_batch_create(user)
            return
        mail_batch_match = MAIL_BATCH_API_RE.fullmatch(path)
        if mail_batch_match and mail_batch_match.group(2) == "retry":
            self.handle_mail_batch_retry(user, mail_batch_match.group(1))
            return
        if path == "/api/projects":
            self.handle_project_create(user)
            return
        if path == "/api/projects/import":
            self.handle_project_import(user)
            return
        project_match = self.project_api_match(path)
        if project_match:
            project_id, action, file_id = project_match.groups()
            if action == "archive" and file_id is None:
                self.handle_project_archive(user, project_id)
                return
            if action == "files" and file_id is None:
                self.handle_project_file_upload(user, project_id)
                return
        if path == "/api/generate-report":
            self.handle_generate_report()
            return
        if path == "/api/import-bpm-sources":
            self.handle_import_bpm_sources()
            return
        if path in {"/api/bpm-jobs", "/api/bpm-topic-jobs"}:
            self.handle_bpm_job("topic")
            return
        if path == "/api/bpm-author-jobs":
            self.handle_bpm_job("author")
            return
        if path == "/api/bpm-submit":
            self.handle_bpm_submit()
            return
        if path != "/api/export-docx":
            self.send_error(404)
            return
        try:
            payload = self.read_json()
            project_id = payload.get("projectId")
            project = None
            export_payload = payload
            if project_id:
                project = APP_STORE.get_project(user["id"], str(project_id))
                if project is None or project.get("archivedAt"):
                    self.send_json(
                        404,
                        {"error": "未找到选题项目。", "code": "PROJECT_NOT_FOUND"},
                    )
                    return
                export_payload = project_report_payload(project, user)
            else:
                export_payload["editorName"] = user["display_name"]
                export_payload["createdBy"] = user["display_name"]
            out = build_docx(export_payload)
            data = out.read_bytes()
            filename = out.name
            if project is not None:
                record = self.save_project_docx(
                    user,
                    project["id"],
                    "exported_report",
                    filename,
                    data,
                )
                try:
                    latest = APP_STORE.get_project(user["id"], project["id"])
                    state = normalize_project_state(latest["state"])
                    state["docxExported"] = True
                    state["sourceFiles"]["exportedReport"] = self.project_source_file(record)
                    summary = summarize_project_state(state, user["display_name"])
                    APP_STORE.update_project(
                        user["id"],
                        project["id"],
                        latest["version"],
                        {
                            "title": summary["title"],
                            "author_name": summary["authorName"],
                            "editor_name": summary["editorName"],
                            "report_status": summary["reportStatus"],
                            "author_status": summary["authorStatus"],
                            "bpm_status": summary["bpmStatus"],
                            "state": summary["state"],
                        },
                        revision_reason="export_docx",
                    )
                except Exception:
                    self.rollback_project_file(user, project["id"], record)
                    raise
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            self.send_header("Content-Length", str(len(data)))
            self.send_header("Content-Disposition", f"attachment; filename*=UTF-8''{quote(filename)}")
            self.end_headers()
            self.wfile.write(data)
        except Exception as exc:
            body = str(exc).encode("utf-8")
            self.send_response(500)
            self.send_header("Content-Type", "text/plain; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

    def handle_bpm_job(self, job_type: str):
        try:
            user = self.require_user()
            if not user:
                return
            length = int(self.headers.get("Content-Length", "0"))
            payload = json.loads(self.rfile.read(length).decode("utf-8"))
            result = create_bpm_job(payload, user, job_type)
            body = json.dumps(result, ensure_ascii=False).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
        except Exception as exc:
            body = str(exc).encode("utf-8")
            self.send_response(500)
            self.send_header("Content-Type", "text/plain; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

    def handle_bpm_submit(self):
        try:
            user = self.require_user()
            if not user:
                return
            length = int(self.headers.get("Content-Length", "0"))
            payload = json.loads(self.rfile.read(length).decode("utf-8"))
            trusted_payload, _ = trusted_bpm_payload(payload, user)
            result = run_bpm_submit(trusted_payload)
            body = json.dumps(result, ensure_ascii=False).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
        except Exception as exc:
            body = str(exc).encode("utf-8")
            self.send_response(500)
            self.send_header("Content-Type", "text/plain; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

    def handle_generate_report(self):
        try:
            runtime_log(f"generate-report start from={self.client_address[0]} length={self.headers.get('Content-Length', '0')}")
            form = read_multipart_form(
                fp=self.rfile,
                headers=self.headers,
                environ={
                    "REQUEST_METHOD": "POST",
                    "CONTENT_TYPE": self.headers.get("Content-Type", ""),
                    "CONTENT_LENGTH": self.headers.get("Content-Length", "0"),
                },
            )
            file_item = form["file"] if "file" in form else None
            if file_item is None or not getattr(file_item, "file", None):
                raise ValueError("请上传选题申报表 DOCX 文件。")
            file_bytes = file_item.file.read()
            filename = getattr(file_item, "filename", "application.docx")
            project_id = form.getfirst("projectId", "")
            if isinstance(project_id, bytes):
                project_id = project_id.decode("utf-8")
            project = None
            user = self.current_user()
            expected_version = None
            if project_id:
                project = APP_STORE.get_project(user["id"], str(project_id)) if user else None
                if project is None or project.get("archivedAt"):
                    self.send_json(
                        404,
                        {"error": "未找到选题项目。", "code": "PROJECT_NOT_FOUND"},
                    )
                    return
                raw_version = form.getfirst("version", "")
                if isinstance(raw_version, bytes):
                    raw_version = raw_version.decode("utf-8")
                expected_version = int(raw_version) if str(raw_version).strip() else project["version"]
            runtime_log(f"generate-report file filename={filename} bytes={len(file_bytes)}")
            result = generate_report_from_upload(file_bytes, filename)
            response = dict(result)
            if project is not None:
                record = self.save_project_docx(
                    user,
                    project["id"],
                    "application",
                    filename,
                    file_bytes,
                )
                try:
                    project = self.persist_project_report(
                        user,
                        project,
                        result,
                        source_kind="generated",
                        source_files={
                            "application": self.project_source_file(record),
                        },
                        expected_version=expected_version,
                        revision_reason="generated_report",
                    )
                except Exception:
                    self.rollback_project_file(user, project["id"], record)
                    raise
                response.update(self.project_payload(user, project, include_preflight=True))
            body = json.dumps(response, ensure_ascii=False).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            runtime_log(f"generate-report ok title={result.get('title', '')} bytes={len(body)}")
        except ModelServiceError as exc:
            runtime_log(f"generate-report model error type={type(exc).__name__}")
            body = str(exc).encode("utf-8")
            self.send_response(503)
            self.send_header("Content-Type", "text/plain; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
        except Exception as exc:
            runtime_log(f"generate-report error {type(exc).__name__}: {exc}")
            body = str(exc).encode("utf-8")
            self.send_response(500)
            self.send_header("Content-Type", "text/plain; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

    def handle_import_bpm_sources(self):
        try:
            runtime_log(f"import-bpm-sources start from={self.client_address[0]} length={self.headers.get('Content-Length', '0')}")
            form = read_multipart_form(
                fp=self.rfile,
                headers=self.headers,
                environ={
                    "REQUEST_METHOD": "POST",
                    "CONTENT_TYPE": self.headers.get("Content-Type", ""),
                    "CONTENT_LENGTH": self.headers.get("Content-Length", "0"),
                },
            )
            application_item = form["application"] if "application" in form else None
            report_item = form["report"] if "report" in form else None
            if application_item is None or not getattr(application_item, "file", None):
                raise ValueError("请上传选题申报表 DOCX。")
            if report_item is None or not getattr(report_item, "file", None):
                raise ValueError("请上传选题策划报告 DOCX。")
            application_bytes = application_item.file.read()
            report_bytes = report_item.file.read()
            application_name = getattr(application_item, "filename", "application.docx")
            report_name = getattr(report_item, "filename", "report.docx")
            result = import_bpm_sources(application_bytes, application_name, report_bytes, report_name)
            body = json.dumps(result, ensure_ascii=False).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            runtime_log(f"import-bpm-sources ok title={result.get('title', '')} bytes={len(body)}")
        except Exception as exc:
            runtime_log(f"import-bpm-sources error {type(exc).__name__}: {exc}")
            body = str(exc).encode("utf-8")
            self.send_response(500)
            self.send_header("Content-Type", "text/plain; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)


@contextmanager
def single_instance_lock(lock_path: Path):
    lock_path = Path(lock_path)
    lock_path.parent.mkdir(parents=True, exist_ok=True)
    handle = lock_path.open("a+")
    try:
        try:
            fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError as error:
            raise RuntimeError("服务已在运行，请勿重复启动。") from error
        yield
    finally:
        try:
            fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
        finally:
            handle.close()


def main():
    host = os.environ.get("PHEI_HOST", "127.0.0.1")
    port = int(os.environ.get("PHEI_PORT", "4174"))
    with single_instance_lock(SERVER_LOCK_PATH):
        APP_STORE.mark_interrupted_jobs_failed()
        APP_STORE.mark_interrupted_mail_batches_failed()
        server = ThreadingHTTPServer((host, port), Handler)
        try:
            print(f"Serving http://{host}:{port}/index.html")
            server.serve_forever()
        finally:
            server.server_close()


if __name__ == "__main__":
    main()
