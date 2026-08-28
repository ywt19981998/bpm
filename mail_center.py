"""Utilities for parsing teacher recipient lists and composing mail."""

from __future__ import annotations

import csv
import io
import re
import smtplib
import ssl
from email.message import EmailMessage
from email.utils import formataddr
from pathlib import PurePath
from typing import Any, Iterable

from openpyxl import load_workbook


NAME_HEADERS = {"姓名", "教师姓名", "老师姓名", "name"}
EMAIL_HEADERS = {"邮箱", "电子邮箱", "e-mail", "email"}
EMAIL_RE = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
VARIABLE_RE = re.compile(r"\{\{\s*(.*?)\s*\}\}")
SMTP_SECURITY_OPTIONS = {"ssl", "starttls", "none"}
SMTP_TIMEOUT_SECONDS = 20


def normalize_header(value: Any) -> str:
    """Normalize a column or variable name for matching."""
    return str(value or "").strip().lower()


def render_mail_text(template: str, variables: dict) -> str:
    normalized = {
        normalize_header(key): str(value or "") for key, value in variables.items()
    }
    return VARIABLE_RE.sub(
        lambda match: normalized.get(normalize_header(match.group(1)), ""),
        str(template or ""),
    )


def validate_mail_compose(
    subject: str, body: str, recipients: list[dict]
) -> list[str]:
    errors = []
    if not str(subject or "").strip():
        errors.append("邮件主题不能为空")
    if not str(body or "").strip():
        errors.append("邮件正文不能为空")
    if not recipients:
        errors.append("至少需要一位收件人")
    for recipient in recipients:
        name = str(recipient.get("name") or "").strip()
        email = str(recipient.get("email") or "").strip()
        if not name:
            errors.append("收件人姓名不能为空")
        if not email:
            errors.append("收件人邮箱不能为空")
        elif not EMAIL_RE.fullmatch(email):
            errors.append("收件人邮箱格式无效")
    return errors


def normalize_smtp_settings(payload: dict) -> dict:
    if not isinstance(payload, dict):
        raise ValueError("SMTP 配置必须是对象")

    host = str(payload.get("host") or "").strip()
    account = str(payload.get("account") or "").strip()
    password = str(payload.get("password") or "")
    from_name = str(payload.get("fromName") or "").strip()
    from_address = str(payload.get("fromAddress") or "").strip()
    security = str(payload.get("security") or "").strip().casefold()
    raw_port = payload.get("port")

    if not host:
        raise ValueError("SMTP 主机不能为空")
    if not account:
        raise ValueError("SMTP 账号不能为空")
    if not password:
        raise ValueError("SMTP 密码或授权码不能为空")
    if not from_name:
        raise ValueError("发件人名称不能为空")
    if not EMAIL_RE.fullmatch(from_address):
        raise ValueError("发件人邮箱格式无效")
    if security not in SMTP_SECURITY_OPTIONS:
        raise ValueError("SMTP 加密方式无效")
    if isinstance(raw_port, bool):
        raise ValueError("SMTP 端口无效")
    try:
        port = int(str(raw_port).strip())
    except (TypeError, ValueError) as error:
        raise ValueError("SMTP 端口无效") from error
    if not 1 <= port <= 65535:
        raise ValueError("SMTP 端口无效")

    return {
        "host": host,
        "port": port,
        "security": security,
        "account": account,
        "password": password,
        "fromName": from_name,
        "fromAddress": from_address,
    }


def test_smtp_connection(settings: dict, smtp_factory=None) -> None:
    normalized = normalize_smtp_settings(settings)
    client = _open_smtp_connection(normalized, smtp_factory=smtp_factory)
    _close_smtp_connection(client)


def send_smtp_message(settings: dict, recipient: dict, smtp_factory=None) -> None:
    normalized = normalize_smtp_settings(settings)
    if not isinstance(recipient, dict):
        raise ValueError("收件人信息无效")
    errors = validate_mail_compose(
        str(recipient.get("subject") or ""),
        str(recipient.get("body") or ""),
        [recipient],
    )
    if errors:
        raise ValueError(errors[0])

    message = EmailMessage()
    message["From"] = formataddr((normalized["fromName"], normalized["fromAddress"]))
    message["To"] = str(recipient["email"]).strip()
    message["Subject"] = str(recipient["subject"]).strip()
    message.set_content(str(recipient["body"]))

    client = _open_smtp_connection(normalized, smtp_factory=smtp_factory)
    try:
        client.send_message(message)
    except Exception as error:
        raise _sanitized_smtp_error(error, normalized["password"]) from error
    finally:
        _close_smtp_connection(client)


def _open_smtp_connection(settings: dict, smtp_factory=None):
    client = None
    try:
        context = ssl.create_default_context()
        if settings["security"] == "ssl":
            factory = smtp_factory or smtplib.SMTP_SSL
            client = factory(
                settings["host"],
                settings["port"],
                timeout=SMTP_TIMEOUT_SECONDS,
                context=context,
            )
        else:
            factory = smtp_factory or smtplib.SMTP
            client = factory(
                settings["host"], settings["port"], timeout=SMTP_TIMEOUT_SECONDS
            )
            if settings["security"] == "starttls":
                client.starttls(context=context)
        client.login(settings["account"], settings["password"])
        return client
    except Exception as error:
        if client is not None:
            _close_smtp_connection(client)
        raise _sanitized_smtp_error(error, settings["password"]) from error


def _close_smtp_connection(client) -> None:
    try:
        client.quit()
    except Exception:
        try:
            client.close()
        except Exception:
            pass


def _sanitized_smtp_error(error: Exception, password: str) -> RuntimeError:
    summary = str(error).replace(password, "[REDACTED]")
    return RuntimeError(f"SMTP {type(error).__name__}: {summary}")


def parse_recipient_file(data: bytes, filename: str) -> dict:
    columns, rows = _read_rows(data, filename)
    name_column = _find_column(columns, NAME_HEADERS)
    email_column = _find_column(columns, EMAIL_HEADERS)
    if name_column is None or email_column is None:
        raise ValueError("名单必须包含姓名和邮箱表头")

    valid = []
    invalid = []
    duplicates = 0
    seen_emails = set()
    for row_number, row in enumerate(rows, start=2):
        record = _record_from_row(row, columns, name_column, email_column)
        name = record["name"]
        email = record["email"]
        if not name:
            invalid.append(
                {"row": row_number, "name": name, "email": email, "reason": "姓名不能为空"}
            )
            continue
        if not EMAIL_RE.fullmatch(email):
            invalid.append(
                {"row": row_number, "name": name, "email": email, "reason": "邮箱格式无效"}
            )
            continue
        email_key = email.casefold()
        if email_key in seen_emails:
            duplicates += 1
            continue
        seen_emails.add(email_key)
        valid.append(record)

    return {
        "valid": valid,
        "invalid": invalid,
        "duplicates": duplicates,
        "columns": columns,
    }


def _read_rows(
    data: bytes, filename: str
) -> tuple[list[str], list[dict[str, Any]]]:
    suffix = PurePath(filename).suffix.casefold()
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx_rows(data)
    if suffix != ".csv":
        raise ValueError("仅支持 CSV 或 XLSX 文件")
    return _read_csv_rows(data)


def _read_csv_rows(data: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    last_error = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = data.decode(encoding)
            reader = csv.DictReader(io.StringIO(text))
            if reader.fieldnames is None:
                return [], []
            columns = list(reader.fieldnames)
            return columns, [dict(row) for row in reader]
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError("CSV 文件编码无法识别") from last_error


def _read_xlsx_rows(data: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        values = worksheet.iter_rows(values_only=True)
        try:
            header_row = next(values)
        except StopIteration:
            return [], []
        columns = [str(value or "").strip() for value in header_row]
        rows = [
            {column: value for column, value in zip(columns, row)} for row in values
        ]
        return columns, rows
    finally:
        workbook.close()


def _find_column(columns: Iterable[str], aliases: set[str]) -> str | None:
    for column in columns:
        if normalize_header(column) in aliases:
            return column
    return None


def _record_from_row(
    row: dict[str, Any], columns: list[str], name_column: str, email_column: str
) -> dict[str, str]:
    name = "" if row.get(name_column) is None else str(row.get(name_column)).strip()
    email = "" if row.get(email_column) is None else str(row.get(email_column)).strip()
    record = {
        column: "" if row.get(column) is None else str(row.get(column)).strip()
        for column in columns
    }
    record["name"] = name
    record["email"] = email
    return record
