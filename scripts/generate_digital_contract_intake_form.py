from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "output" / "doc" / "数字教材出版合同信息采集表.docx"

CONTRACT_INFO_FIELDS = ["选题名称", "署名方式"]
PARTY_A_FIELDS = ["甲方名称", "著作权人代表", "电话", "身份证号", "邮箱", "通信地址"]
COPYRIGHT_HEADERS = [
    "序号",
    "姓名",
    "证件类型",
    "证件号",
    "著作权人电话",
    "报酬比例(%)",
    "身份证复印件文件名",
    "地址",
    "邮件",
    "邮编",
    "工作单位",
]
DELIVERY_HEADERS = ["序号", "内容类型", "内容明细", "交付数量", "备注"]


def set_run_font(run, name="宋体", size=10.5, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def set_cell_text(cell, text, size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def shade_cell(cell, fill="D9E2F3"):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_width(cell, width_mm):
    cell.width = Mm(width_mm)
    properties = cell._tc.get_or_add_tcPr()
    width = properties.first_child_found_in("w:tcW")
    if width is None:
        width = OxmlElement("w:tcW")
        properties.append(width)
    width.set(qn("w:w"), str(int(width_mm * 56.7)))
    width.set(qn("w:type"), "dxa")


def configure_section(section, landscape=False):
    section.top_margin = Mm(16)
    section.bottom_margin = Mm(16)
    section.left_margin = Mm(15)
    section.right_margin = Mm(15)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Mm(297 if landscape else 210)
    section.page_height = Mm(210 if landscape else 297)


def add_section_heading(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_font(run, name="黑体", size=14, bold=True)


def add_key_value_table(document, fields):
    table = document.add_table(rows=len(fields), cols=2)
    table.style = "Table Grid"
    table.autofit = False
    for row, field in zip(table.rows, fields):
        label_cell, value_cell = row.cells
        set_cell_width(label_cell, 42)
        set_cell_width(value_cell, 138)
        set_cell_text(label_cell, field, bold=True)
        set_cell_text(value_cell, "\u3000" * 12)
        shade_cell(label_cell, "EAF1FB")
        row.height = Mm(11 if field != "通信地址" else 20)
    document.add_paragraph()
    return table


def add_grid_table(document, headers, row_count, widths_mm):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for cell, text, width in zip(table.rows[0].cells, headers, widths_mm):
        set_cell_width(cell, width)
        set_cell_text(cell, text, size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_cell(cell)

    for number in range(1, row_count + 1):
        row = table.add_row()
        row.height = Mm(12)
        for index, (cell, width) in enumerate(zip(row.cells, widths_mm)):
            set_cell_width(cell, width)
            set_cell_text(
                cell,
                str(number) if index == 0 else "\u3000",
                size=9,
                align=WD_ALIGN_PARAGRAPH.CENTER if index in {0, 5} else WD_ALIGN_PARAGRAPH.LEFT,
            )
    return table


def add_delivery_table(document):
    widths_mm = [14, 35, 92, 28, 26]
    table = document.add_table(rows=1, cols=len(DELIVERY_HEADERS))
    table.style = "Table Grid"
    table.autofit = False
    for cell, text, width in zip(table.rows[0].cells, DELIVERY_HEADERS, widths_mm):
        set_cell_width(cell, width)
        set_cell_text(cell, text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_cell(cell)

    for number in range(1, 9):
        row = table.add_row()
        row.height = Mm(15)
        for index, (cell, width) in enumerate(zip(row.cells, widths_mm)):
            set_cell_width(cell, width)
            set_cell_text(
                cell,
                str(number) if index == 0 else "\u3000",
                align=WD_ALIGN_PARAGRAPH.CENTER if index in {0, 3} else WD_ALIGN_PARAGRAPH.LEFT,
            )
    return table


def build_document():
    document = Document()
    configure_section(document.sections[0])
    normal_style = document.styles["Normal"]
    normal_style.font.name = "宋体"
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal_style.font.size = Pt(10.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    title_run = title.add_run("数字教材出版合同信息采集表")
    set_run_font(title_run, name="黑体", size=18, bold=True)

    notice = document.add_paragraph()
    notice.paragraph_format.space_after = Pt(8)
    notice_run = notice.add_run(
        "填写说明：请如实填写以下信息，并将身份证复印件、授权委托书、著作权证明书等材料以独立文件形式与本表一并提交。"
    )
    set_run_font(notice_run, size=10.5)

    add_section_heading(document, "一、合同信息")
    add_key_value_table(document, CONTRACT_INFO_FIELDS)

    add_section_heading(document, "二、甲方信息")
    add_key_value_table(document, PARTY_A_FIELDS)

    contributor_section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section(contributor_section, landscape=True)
    add_section_heading(document, "三、全体著作权人表")
    instruction = document.add_paragraph()
    instruction_run = instruction.add_run("请填写全部著作权人信息；报酬比例合计应为 100%。")
    set_run_font(instruction_run, size=10.5)
    add_grid_table(
        document,
        COPYRIGHT_HEADERS,
        row_count=10,
        widths_mm=[12, 20, 20, 29, 27, 20, 30, 37, 30, 17, 27],
    )

    financial_section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section(financial_section)
    add_section_heading(document, "四、出版资助")
    add_key_value_table(document, ["是否有出版资助（是/否）", "出版资助金额（元）", "资助支付截止日期"])
    financial_note = document.add_paragraph()
    note_run = financial_note.add_run("如无出版资助，请在第一项填写“否”，其余两项留空。")
    set_run_font(note_run, size=10.5)

    add_section_heading(document, "五、交付清单")
    delivery_note = document.add_paragraph()
    delivery_run = delivery_note.add_run("请按实际交付内容填写；本表将作为合同附件“甲方应交付作品内容清单”的数据来源。")
    set_run_font(delivery_run, size=10.5)
    add_delivery_table(document)

    document.add_paragraph()
    attachment_note = document.add_paragraph()
    attachment_run = attachment_note.add_run("随表提交材料：□ 身份证复印件  □ 授权委托书  □ 著作权证明书  □ 其他：________________")
    set_run_font(attachment_run, size=10.5)
    return document


def main():
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    build_document().save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
