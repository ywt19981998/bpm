import unittest
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Mm


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "output" / "doc" / "数字教材出版合同信息采集表.docx"
COPYRIGHT_HEADERS = [
    "序号",
    "姓名",
    "证件类型",
    "证件号",
    "著作权人电话",
    "报酬比例(%)",
    "身份证复印件文件名",
    "地址",
    "邮件",
    "邮编",
    "工作单位",
]
DELIVERY_HEADERS = ["序号", "内容类型", "内容明细", "交付数量", "备注"]


def header_cells(table):
    return [cell.text.strip() for cell in table.rows[0].cells]


class DigitalContractIntakeFormTest(unittest.TestCase):
    def test_generated_form_has_approved_sections_and_tables(self):
        self.assertTrue(OUTPUT_PATH.exists(), "Run the intake-form generator first.")

        document = Document(OUTPUT_PATH)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)

        for heading in [
            "一、合同信息",
            "二、甲方信息",
            "三、全体著作权人表",
            "四、出版资助",
            "五、交付清单",
        ]:
            self.assertIn(heading, text)

        self.assertEqual(header_cells(document.tables[2]), COPYRIGHT_HEADERS)
        self.assertEqual(header_cells(document.tables[4]), DELIVERY_HEADERS)
        self.assertEqual(len(document.tables[2].rows), 11)
        self.assertEqual(len(document.tables[4].rows), 9)
        self.assertLessEqual(document.tables[2].rows[1].height, Mm(12))
        self.assertTrue(any(section.orientation == WD_ORIENT.LANDSCAPE for section in document.sections))


if __name__ == "__main__":
    unittest.main()
