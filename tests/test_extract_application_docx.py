import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
EXTRACTOR_PATH = (
    ROOT
    / "skills"
    / "phei-topic-planning-report"
    / "scripts"
    / "extract_application_docx.py"
)
SPEC = importlib.util.spec_from_file_location("test_extract_application_docx_module", EXTRACTOR_PATH)
extractor = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(extractor)


class ExtractApplicationDocxTests(unittest.TestCase):
    def test_canonical_key_normalizes_spaced_name_and_topic_alias(self):
        self.assertEqual(extractor.canonical_key("姓 名"), "姓名")
        self.assertEqual(extractor.canonical_key("选题名"), "选题名称")
        self.assertEqual(extractor.canonical_key("选题名称（暂定）"), "选题名称")

    def test_build_payload_reads_known_label_after_section_cell(self):
        document = Document()
        topic_table = document.add_table(rows=1, cols=3)
        topic_table.rows[0].cells[0].text = "选 题 情 况"
        topic_table.rows[0].cells[1].text = "选题名"
        topic_table.rows[0].cells[2].text = "数智赋能：西门子S7-1200 PLC项目式教程"

        author_table = document.add_table(rows=1, cols=2)
        author_table.rows[0].cells[0].text = "姓 名"
        author_table.rows[0].cells[1].text = "朱万浩"

        with tempfile.NamedTemporaryFile(suffix=".docx") as tmp:
            document.save(tmp.name)
            payload = extractor.build_payload(Path(tmp.name), include_sensitive=False)

        self.assertEqual(
            payload["canonical_fields"]["选题名称"],
            "数智赋能：西门子S7-1200 PLC项目式教程",
        )
        self.assertEqual(payload["canonical_fields"]["姓名"], "朱万浩")

    def test_non_sensitive_payload_masks_raw_table_values(self):
        document = Document()
        table = document.add_table(rows=2, cols=4)
        table.rows[0].cells[0].text = "姓名"
        table.rows[0].cells[1].text = "测试作者"
        table.rows[0].cells[2].text = "联系电话"
        table.rows[0].cells[3].text = "13800138000"
        table.rows[1].cells[0].text = "身份证号"
        table.rows[1].cells[1].text = "110101199001011234"
        table.rows[1].cells[2].text = "电子邮箱"
        table.rows[1].cells[3].text = "private@example.com"

        with tempfile.NamedTemporaryFile(suffix=".docx") as tmp:
            document.save(tmp.name)
            payload = extractor.build_payload(Path(tmp.name), include_sensitive=False)

        serialized = json.dumps(payload, ensure_ascii=False)
        self.assertNotIn("13800138000", serialized)
        self.assertNotIn("110101199001011234", serialized)
        self.assertNotIn("private@example.com", serialized)
        self.assertIn("[已隐藏]", serialized)

    def test_private_payload_canonicalizes_author_identity_and_contact_fields(self):
        document = Document()
        table = document.add_table(rows=4, cols=8)
        values = [
            ["姓 名", "王明", "性 别", "男", "职 称", "正高级工程师", "学 历", "博士"],
            ["邮政编码", "100000", "E-mail", "author@example.com", "", "", "", ""],
            ["电 话", "(O)", "(H)", "(手机) 13800138000", "", "", "", ""],
            ["个人简历(学习经历和工作经历)", "作者工作简历", "", "", "", "", "", ""],
        ]
        for row, row_values in zip(table.rows, values):
            for cell, value in zip(row.cells, row_values):
                cell.text = value

        with tempfile.NamedTemporaryFile(suffix=".docx") as tmp:
            document.save(tmp.name)
            payload = extractor.build_payload(Path(tmp.name), include_sensitive=True)
            safe_payload = extractor.build_payload(Path(tmp.name), include_sensitive=False)

        fields = payload["canonical_fields"]
        self.assertEqual(fields["姓名"], "王明")
        self.assertEqual(fields["性别"], "男")
        self.assertEqual(fields["职称"], "正高级工程师")
        self.assertEqual(fields["学历"], "博士")
        self.assertEqual(fields["电子邮箱"], "author@example.com")
        self.assertEqual(fields["电话"], "13800138000")
        self.assertEqual(fields["个人简历"], "作者工作简历")
        self.assertNotIn("13800138000", json.dumps(safe_payload, ensure_ascii=False))


if __name__ == "__main__":
    unittest.main()
