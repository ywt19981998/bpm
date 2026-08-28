import io
import unittest
import zipfile
from unittest.mock import patch

from openpyxl import Workbook

from mail_center import (
    normalize_smtp_settings,
    parse_recipient_file,
    render_mail_text,
    send_smtp_message,
    test_smtp_connection,
    validate_mail_compose,
)


class MailRecipientTests(unittest.TestCase):
    def test_csv_aliases_and_duplicate_email(self):
        data = "教师姓名,电子邮箱,学校\n张三,ZHANG@example.com,A大学\n李四,zhang@example.com,B大学\n".encode("utf-8-sig")

        result = parse_recipient_file(data, "teachers.csv")

        self.assertEqual(len(result["valid"]), 1)
        self.assertEqual(result["duplicates"], 1)
        self.assertEqual(result["valid"][0]["name"], "张三")
        self.assertEqual(result["valid"][0]["email"], "ZHANG@example.com")
        self.assertEqual(result["valid"][0]["学校"], "A大学")
        self.assertEqual(result["columns"], ["教师姓名", "电子邮箱", "学校"])

    def test_xlsx_rows_are_parsed_from_an_in_memory_workbook(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["姓名", "邮箱", "学院"])
        sheet.append(["王五", "wang@example.com", "信息学院"])
        buffer = io.BytesIO()
        workbook.save(buffer)

        result = parse_recipient_file(buffer.getvalue(), "teachers.xlsx")

        self.assertEqual(result["valid"][0]["name"], "王五")
        self.assertEqual(result["valid"][0]["email"], "wang@example.com")
        self.assertEqual(result["valid"][0]["姓名"], "王五")
        self.assertEqual(result["valid"][0]["邮箱"], "wang@example.com")
        self.assertEqual(result["valid"][0]["学院"], "信息学院")
        self.assertEqual(result["invalid"], [])
        self.assertEqual(result["duplicates"], 0)

    def test_invalid_email_is_reported_without_becoming_valid(self):
        data = "姓名,邮箱\n张三,not-an-email\n李四,li@example.com\n".encode("utf-8-sig")

        result = parse_recipient_file(data, "teachers.csv")

        self.assertEqual([item["name"] for item in result["valid"]], ["李四"])
        self.assertEqual(len(result["invalid"]), 1)
        self.assertEqual(result["invalid"][0]["row"], 2)
        self.assertIn("邮箱", result["invalid"][0]["reason"])

    def test_missing_required_headers_is_rejected(self):
        data = "教师,单位\n张三,A大学\n".encode("utf-8-sig")

        with self.assertRaises(ValueError):
            parse_recipient_file(data, "teachers.csv")

    def test_gb18030_csv_is_supported(self):
        data = "姓名,邮箱\n张三,zhang@example.com\n".encode("gb18030")

        result = parse_recipient_file(data, "teachers.csv")

        self.assertEqual(result["valid"][0]["name"], "张三")

    def test_missing_optional_variable_becomes_empty(self):
        rendered = render_mail_text("{{姓名}}老师，您好，{{学院}}", {"姓名": "王五"})

        self.assertEqual(rendered, "王五老师，您好，")

    def test_variable_names_are_trimmed_and_values_are_stringified(self):
        rendered = render_mail_text("{{ 姓名 }}：{{职称}}", {"姓名": "王五", "职称": 3})

        self.assertEqual(rendered, "王五：3")

    def test_compose_validation_reports_required_fields(self):
        errors = validate_mail_compose("  ", "", [])

        self.assertEqual(errors, ["邮件主题不能为空", "邮件正文不能为空", "至少需要一位收件人"])

    def test_compose_validation_accepts_a_complete_message(self):
        errors = validate_mail_compose(
            "通知",
            "{{姓名}}老师，您好",
            [{"name": "张三", "email": "zhang@example.com"}],
        )

        self.assertEqual(errors, [])

    def test_compose_validation_rejects_recipient_without_name(self):
        errors = validate_mail_compose(
            "通知",
            "正文",
            [{"email": "zhang@example.com"}],
        )

        self.assertIn("收件人姓名不能为空", errors)

    def test_compose_validation_rejects_recipient_without_email(self):
        errors = validate_mail_compose(
            "通知",
            "正文",
            [{"name": "张三"}],
        )

        self.assertIn("收件人邮箱不能为空", errors)

    def test_compose_validation_rejects_recipient_with_invalid_email(self):
        errors = validate_mail_compose(
            "通知",
            "正文",
            [{"name": "张三", "email": "not-an-email"}],
        )

        self.assertIn("收件人邮箱格式无效", errors)

    def test_empty_csv_with_required_headers_is_valid(self):
        data = "姓名,邮箱\n".encode("utf-8-sig")

        result = parse_recipient_file(data, "teachers.csv")

        self.assertEqual(result["valid"], [])
        self.assertEqual(result["invalid"], [])
        self.assertEqual(result["columns"], ["姓名", "邮箱"])

    def test_empty_xlsx_with_required_headers_is_valid(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["姓名", "邮箱"])
        buffer = io.BytesIO()
        workbook.save(buffer)

        result = parse_recipient_file(buffer.getvalue(), "teachers.xlsx")

        self.assertEqual(result["valid"], [])
        self.assertEqual(result["invalid"], [])
        self.assertEqual(result["columns"], ["姓名", "邮箱"])

    def test_recipient_file_rejects_more_than_the_send_limit(self):
        rows = ["姓名,邮箱"] + [
            f"老师{index},teacher{index}@example.edu.cn" for index in range(1001)
        ]

        with self.assertRaisesRegex(ValueError, "1000"):
            parse_recipient_file("\n".join(rows).encode("utf-8"), "teachers.csv")

    def test_xlsx_rejects_excessive_uncompressed_content(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["姓名", "邮箱"])
        buffer = io.BytesIO()
        workbook.save(buffer)

        expanded = io.BytesIO(buffer.getvalue())
        with zipfile.ZipFile(expanded, "a", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("xl/oversized.bin", b"x" * 2048)

        with patch("mail_center.MAX_XLSX_UNCOMPRESSED_BYTES", 1024):
            with self.assertRaisesRegex(ValueError, "解压"):
                parse_recipient_file(expanded.getvalue(), "teachers.xlsx")


class FakeSMTP:
    def __init__(self):
        self.login_args = None
        self.messages = []
        self.starttls_calls = 0
        self.quit_calls = 0

    def starttls(self, *, context):
        self.starttls_calls += 1
        self.starttls_context = context

    def login(self, account, password):
        self.login_args = (account, password)

    def send_message(self, message):
        self.messages.append(message)

    def quit(self):
        self.quit_calls += 1


class FailingLoginSMTP(FakeSMTP):
    def login(self, account, password):
        raise RuntimeError(f"authentication rejected: {password}")


class MailSmtpTests(unittest.TestCase):
    @staticmethod
    def smtp_settings(**overrides):
        settings = {
            "host": "smtp.example.com",
            "port": 465,
            "security": "ssl",
            "account": "editor@example.com",
            "password": "smtp-auth-code",
            "fromName": "张编辑",
            "fromAddress": "editor@example.com",
        }
        settings.update(overrides)
        return settings

    def test_send_uses_one_recipient_and_never_exposes_password(self):
        fake = FakeSMTP()

        send_smtp_message(
            self.smtp_settings(),
            {
                "name": "张三",
                "email": "teacher@example.edu.cn",
                "subject": "教材合作",
                "body": "张老师您好",
            },
            smtp_factory=lambda *_args, **_kwargs: fake,
        )

        self.assertEqual(fake.login_args, ("editor@example.com", "smtp-auth-code"))
        self.assertEqual(fake.messages[0]["To"], "teacher@example.edu.cn")
        self.assertEqual(fake.messages[0]["Subject"], "教材合作")
        self.assertIn("张老师您好", fake.messages[0].get_content())
        self.assertNotIn("smtp-auth-code", fake.messages[0].as_string())
        self.assertEqual(fake.quit_calls, 1)

    def test_connection_supports_ssl_starttls_and_plain_transport(self):
        ssl_client = FakeSMTP()
        starttls_client = FakeSMTP()
        plain_client = FakeSMTP()

        test_smtp_connection(
            self.smtp_settings(security="ssl"),
            smtp_factory=lambda *_args, **_kwargs: ssl_client,
        )
        test_smtp_connection(
            self.smtp_settings(port=587, security="starttls"),
            smtp_factory=lambda *_args, **_kwargs: starttls_client,
        )
        test_smtp_connection(
            self.smtp_settings(port=25, security="none"),
            smtp_factory=lambda *_args, **_kwargs: plain_client,
        )

        self.assertEqual(ssl_client.starttls_calls, 0)
        self.assertEqual(starttls_client.starttls_calls, 1)
        self.assertEqual(plain_client.starttls_calls, 0)
        self.assertEqual(ssl_client.quit_calls, 1)
        self.assertEqual(starttls_client.quit_calls, 1)
        self.assertEqual(plain_client.quit_calls, 1)

    def test_normalize_rejects_invalid_port_and_security(self):
        with self.assertRaisesRegex(ValueError, "端口"):
            normalize_smtp_settings(self.smtp_settings(port=0))
        with self.assertRaisesRegex(ValueError, "加密"):
            normalize_smtp_settings(self.smtp_settings(security="tls-unknown"))

    def test_login_error_is_sanitized(self):
        with self.assertRaisesRegex(RuntimeError, "authentication rejected") as context:
            test_smtp_connection(
                self.smtp_settings(),
                smtp_factory=lambda *_args, **_kwargs: FailingLoginSMTP(),
            )

        self.assertNotIn("smtp-auth-code", str(context.exception))

if __name__ == "__main__":
    unittest.main()
