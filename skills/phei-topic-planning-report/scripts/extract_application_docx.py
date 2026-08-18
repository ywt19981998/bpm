#!/usr/bin/env python3
"""Extract structured facts from a PHEI topic application DOCX."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Iterable


SENSITIVE_KEYWORDS = {
    "身份证",
    "通信地址",
    "电话",
    "电 话",
    "e-mail",
    "email",
    "邮箱",
    "邮政编码",
}

FIELD_KEY_ALIASES = {
    "姓名": "姓名",
    "作者姓名": "作者姓名",
    "第一作者姓名": "第一作者姓名",
    "主要作（译）者姓名": "主要作（译）者姓名",
    "主要作译者姓名": "主要作（译）者姓名",
    "选题名": "选题名称",
    "选题名称": "选题名称",
    "教材名称": "教材名称",
    "书名": "书名",
    "联系电话": "联系电话",
    "联系电话1": "联系电话1",
    "联系电话2": "联系电话2",
    "电话": "电话",
    "手机": "手机",
    "身份证": "身份证号",
    "身份证号": "身份证号",
    "证件号": "证件号",
    "电子邮箱": "电子邮箱",
    "邮箱": "邮箱",
    "email": "email",
    "e-mail": "电子邮箱",
    "邮政编码": "邮政编码",
    "邮编": "邮编",
    "通信地址": "通信地址",
    "通讯地址": "通讯地址",
    "性别": "性别",
    "职称": "职称",
    "职务": "职务",
    "学历": "学历",
    "学位": "学位",
    "专业": "专业",
    "毕业院校、专业和时间": "毕业院校",
    "单位名称": "单位名称",
    "工作单位": "工作单位",
    "从事方向": "著作方向",
    "个人简历": "个人简历",
    "参加的学术组织及任职": "参加的学术组织及任职情况",
    "参加的学术组织及任职情况": "参加的学术组织及任职情况",
    "科研或教研项目经历": "科研或教研项目经历",
    "所承担过的重点科研或教研项目以及在项目中所承担的工作": "科研或教研项目经历",
    "教学成果获奖情况、作品获奖情况": "获奖情况",
    "主要著作出版情况": "主要著作出版情况",
}

KNOWN_FIELD_KEYS = set(FIELD_KEY_ALIASES.values())


def load_document(path: Path):
    try:
        from docx import Document
    except ImportError as exc:
        raise SystemExit(
            "Missing python-docx. Run this script with the Codex bundled Python "
            "or install python-docx in the active environment."
        ) from exc
    return Document(str(path))


def normalize(text: str) -> str:
    text = (text or "").replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def collapse_repeated(values: Iterable[str]) -> list[str]:
    result: list[str] = []
    for value in values:
        value = normalize(value)
        if value and (not result or result[-1] != value):
            result.append(value)
    return result


def is_sensitive_key(key: str) -> bool:
    key_l = re.sub(r"\s+", "", key).lower()
    return any(word.lower() in key_l for word in SENSITIVE_KEYWORDS)


def canonical_key(key: str) -> str:
    key = normalize(key)
    key = re.sub(r"（[^）]*）", "", key)
    key = re.sub(r"\([^)]*\)", "", key)
    key = normalize(key)
    compact = re.sub(r"\s+", "", key)
    return FIELD_KEY_ALIASES.get(compact.lower(), key)


def mask_value(value: str) -> str:
    if not value:
        return value
    return "[已隐藏]"


def mask_sensitive_patterns(value: str) -> str:
    text = str(value or "")
    text = re.sub(r"(?i)\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", "[已隐藏]", text)
    text = re.sub(r"(?<!\d)1[3-9]\d{9}(?!\d)", "[已隐藏]", text)
    return re.sub(r"(?<!\d)\d{17}[\dXx](?!\d)", "[已隐藏]", text)


def extract_table_rows(doc) -> list[dict]:
    rows: list[dict] = []
    for table_index, table in enumerate(doc.tables):
        for row_index, row in enumerate(table.rows):
            cells = collapse_repeated(cell.text for cell in row.cells)
            if cells:
                rows.append(
                    {
                        "table_index": table_index,
                        "row_index": row_index,
                        "cells": cells,
                    }
                )
    return rows


def mask_table_rows(table_rows: list[dict]) -> list[dict]:
    masked_rows: list[dict] = []
    for row in table_rows:
        cells = list(row.get("cells") or [])
        for index, cell in enumerate(cells):
            if not is_sensitive_key(canonical_key(cell)):
                continue
            next_index = index + 1
            while next_index < len(cells):
                next_cell = cells[next_index]
                next_key = canonical_key(next_cell)
                if next_key in KNOWN_FIELD_KEYS or is_sensitive_key(next_key):
                    break
                cells[next_index] = mask_value(next_cell)
                next_index += 1
            if re.search(r"[:：]", cell):
                label = re.split(r"[:：]", cell, maxsplit=1)[0]
                cells[index] = f"{label}：[已隐藏]"
        masked_rows.append({**row, "cells": cells})
    return masked_rows


def add_field(fields: dict[str, str], key: str, value: str, include_sensitive: bool) -> None:
    key = normalize(key)
    value = normalize(value)
    if not key or not value or key == value:
        return
    if is_sensitive_key(canonical_key(key)) and not include_sensitive:
        value = mask_value(value)
    elif not include_sensitive:
        value = mask_sensitive_patterns(value)
    if key in fields and fields[key] != value:
        fields[key] = fields[key] + "\n" + value
    else:
        fields[key] = value


def extract_field_pairs(table_rows: list[dict], include_sensitive: bool) -> dict[str, str]:
    fields: dict[str, str] = {}
    for row in table_rows:
        cells = row["cells"]
        if len(cells) < 2:
            continue

        # A merged section heading may occupy the first cell, leaving a real
        # label/value pair at offsets 1 and 2. Scan recognized labels before
        # retaining the historical fixed-pair behavior below.
        for index, key in enumerate(cells[:-1]):
            value = cells[index + 1]
            simple_key = canonical_key(key)
            if simple_key not in KNOWN_FIELD_KEYS:
                continue
            if canonical_key(value) in KNOWN_FIELD_KEYS:
                continue
            add_field(fields, key, value, include_sensitive)

        if len(cells) == 2:
            pairs = [(cells[0], cells[1])]
        else:
            pairs = []
            i = 0
            while i + 1 < len(cells):
                key, value = cells[i], cells[i + 1]
                if key and value and len(key) <= 40:
                    pairs.append((key, value))
                i += 2

        for key, value in pairs:
            if canonical_key(value) in KNOWN_FIELD_KEYS:
                continue
            add_field(fields, key, value, include_sensitive)
    return fields


def recover_multi_cell_contact_fields(
    table_rows: list[dict], fields: dict[str, str], include_sensitive: bool
) -> None:
    for row in table_rows:
        cells = list(row.get("cells") or [])
        if not cells or canonical_key(cells[0]) != "电话":
            continue
        mobile = re.search(r"(?<!\d)1[3-9]\d{9}(?!\d)", " ".join(cells[1:]))
        if not mobile:
            continue
        for key in list(fields):
            if canonical_key(key) == "电话":
                fields.pop(key, None)
        fields["电话"] = mobile.group(0) if include_sensitive else mask_value(mobile.group(0))
        return


def extract_paragraphs(doc) -> list[str]:
    return [normalize(p.text) for p in doc.paragraphs if normalize(p.text)]


def extract_heading_section(paragraphs: list[str], heading_patterns: list[str]) -> str:
    start = None
    for i, para in enumerate(paragraphs):
        compact = para.replace(" ", "")
        if any(compact == pattern.replace(" ", "") for pattern in heading_patterns):
            start = i + 1
            break
    if start is None:
        return ""

    stop_headings = [
        "适合读者",
        "读者定位",
        "电子工业出版社选题申报表",
        "大纲及目录",
        "大 纲 及 目 录",
    ]
    collected: list[str] = []
    for para in paragraphs[start:]:
        if collected and any(h in para for h in stop_headings):
            break
        collected.append(para)
    return "\n".join(collected).strip()


def extract_date(paragraphs: list[str]) -> str:
    for para in paragraphs:
        if "填表日期" in para:
            return para
    return ""


def build_canonical_fields(fields: dict[str, str]) -> dict[str, str]:
    canonical: dict[str, str] = {}
    for key, value in fields.items():
        simple = canonical_key(key)
        if not simple:
            continue
        if simple in canonical and canonical[simple] != value:
            canonical[simple] = canonical[simple] + "\n" + value
        else:
            canonical[simple] = value
    return canonical


def build_payload(path: Path, include_sensitive: bool) -> dict:
    doc = load_document(path)
    paragraphs = extract_paragraphs(doc)
    raw_table_rows = extract_table_rows(doc)
    safe_table_rows = mask_table_rows(raw_table_rows)
    fields = extract_field_pairs(raw_table_rows, include_sensitive=include_sensitive)
    recover_multi_cell_contact_fields(raw_table_rows, fields, include_sensitive)

    appendix = {
        "outline_from_paragraphs": extract_heading_section(paragraphs, ["大纲及目录", "大 纲 及 目 录"]),
        "detailed_readers_from_paragraphs": extract_heading_section(paragraphs, ["适合读者"]),
    }

    return {
        "source_file": str(path),
        "document_title": paragraphs[0] if paragraphs else "",
        "form_date": extract_date(paragraphs),
        "application_fields": fields,
        "canonical_fields": build_canonical_fields(fields),
        "appendix_sections": appendix,
        "paragraphs": paragraphs,
        "table_rows": raw_table_rows if include_sensitive else safe_table_rows,
        "safe_table_rows": safe_table_rows,
        "notes": [
            "Sensitive fields are masked by default. Re-run with --include-sensitive only if the user explicitly needs raw private data.",
            "Use application_fields as the primary source; use appendix_sections and paragraphs to recover attached outline/reader details.",
        ],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path, help="Filled topic application .docx")
    parser.add_argument("--out", type=Path, help="Write JSON to this path")
    parser.add_argument(
        "--include-sensitive",
        action="store_true",
        help="Do not mask ID, contact, and address-like fields.",
    )
    args = parser.parse_args()

    if not args.docx.exists():
        parser.error(f"File not found: {args.docx}")
    payload = build_payload(args.docx, include_sensitive=args.include_sensitive)
    text = json.dumps(payload, ensure_ascii=False, indent=2)
    if args.out:
        args.out.write_text(text + "\n", encoding="utf-8")
    else:
        print(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
